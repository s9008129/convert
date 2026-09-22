#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
會議紀錄 artifact 驗收檢查器（deterministic acceptance gate；Stage 05 用）。

用途：
    真實 end-to-end 產出流程結束後，對產生的會議紀錄 artifact（Markdown 與／或
    DOCX）做確定性驗收，判斷它是不是「正式會議紀錄」，而不是摘要失敗退回的
    逐字稿 fallback。每個檢查獨立回報 PASS／FAIL／SKIP，最後印出總 verdict；
    任一 FAIL 即 REJECTED（exit 1）。

結構與判定一律由程式碼推導，不寫死猜測值：
    - 章節樣式：backend/core/templates.py 模板註冊表（get_template(None)＝general）。
      artifact 標題命中其他已註冊模板（如科務會議）時，改以該模板的
      required_section_patterns／extra_field_patterns 驗證，與
      backend/services/summarization.py `_validate_summary_quality` 同語意。
    - 模板的 docx_section_pattern（章節白名單）同時作為「章節邊界」來源，
      供決議內容長度計算使用。
    - fallback 偵測：backend/services/task_processor.py 的 SUMMARY_FAILED_BANNER、
      fallback 標題與 backend/api/routes.py 的下載檔名標籤；以 AST 讀取原始碼
      常數（不 import 整套服務鏈，避免 DATA_DIR／依賴副作用，也避開
      backend.services.__init__ 對同名屬性的覆寫）。

使用方式：
    uv run python scripts/e2e/check_record_output.py --md data/outputs/x.md
    uv run python scripts/e2e/check_record_output.py --md x.md --docx x.docx

輸出：
    每個檢查一行：`PASS|FAIL|SKIP <check-id>: <detail>`；最後一行 `VERDICT: ...`。

退出碼：
    0 - ACCEPTED（無任何 FAIL）
    1 - REJECTED（存在 FAIL）
    2 - 參數或環境錯誤（無法推導檢查依據）
"""

from __future__ import annotations

import argparse
import ast
import os
import re
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator, Optional, Sequence

REPO_ROOT = Path(__file__).resolve().parents[2]
TASK_PROCESSOR_SOURCE_PATH = REPO_ROOT / "backend" / "services" / "task_processor.py"
RESULT_ROUTE_SOURCE_PATH = REPO_ROOT / "backend" / "api" / "routes.py"

# 規範明列的 fallback 技術字樣（程式碼推導以外的失效證據）。
FALLBACK_TECHNICAL_MARKERS = ("StableServiceError", "Traceback", "summary_failed")
# SUMMARY_FAILED_BANNER 首行使用的警示符號。
FALLBACK_BANNER_EMOJI = "⚠️"
# 規範明列的待辦關鍵詞（模板未宣告待辦欄位時的備援標記）。
SPEC_ACTION_KEYWORDS = ("待辦事項", "辦理事項", "預定事項")

MIN_BODY_CJK = 600  # 規範：正文至少 600 個中文字
MIN_DECISION_CJK = 20  # 規範：決議區段至少 20 個中文字
MIN_MEANINGFUL_CJK = 100  # 「（待確認）」以外的實質內容下限（防整份只有佔位符）
MIN_SECTION_CJK = 10  # 章節窗口視為「有內容」的下限
MIN_ACTION_CJK = 2  # 待辦／行動項內容下限
MIN_ACTION_TABLE_CJK = 10  # 待辦表格資料列內容下限
MIN_CJK_RATIO = 0.5  # 中文占比下限（其餘為英文／數字）
MIN_DOCX_PARAGRAPHS = 15  # 規範：DOCX 段落數下限
DECISION_WINDOW_CAP = 1200  # 決議內容取樣上限（避免跨章節誤計）
DETAIL_SNIPPET_LEN = 36

CJK_PATTERN = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")
LATIN_PATTERN = re.compile(r"[A-Za-z]")
DIGIT_PATTERN = re.compile(r"[0-9]")
CJK_PHRASE_PATTERN = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]{5,}")
PLACEHOLDER_PATTERN = re.compile(r"[（(]待確認[）)]")
TODO_PATTERN = re.compile(r"TODO")
XXX_PATTERN = re.compile(r"XXX")
DECISION_HEADING_PATTERN = re.compile(r"決議(?:事項)?\s*[:：]")
MARKDOWN_TABLE_ROW_PATTERN = re.compile(r"^\s*\|.*\|\s*$")
MARKDOWN_TABLE_RULE_PATTERN = re.compile(r"^\s*\|[\s:|-]+\|\s*$")


def _bootstrap_backend_import() -> None:
    """設定匯入 backend 所需的最小環境。

    backend.core.logger 於匯入階段就會建立 DATA_DIR/logs，其預設值 /app/data
    在本機不可寫；此處比照 tests/test_asr_subprocess.py，未指定 DATA_DIR 時
    導向暫存目錄，不污染 repo。同時把 repo root 加進 sys.path，讓直接以
    `python scripts/e2e/check_record_output.py` 執行時仍能匯入 backend。
    """
    if str(REPO_ROOT) not in sys.path:
        sys.path.insert(0, str(REPO_ROOT))
    if not os.environ.get("DATA_DIR"):
        os.environ["DATA_DIR"] = tempfile.mkdtemp(prefix="record-output-check-")
    # backend 匯入時會初始化 loguru（主控台 handler 寫 stdout）；驗收報告必須是
    # 乾淨的單行輸出，因此未指定 LOG_LEVEL 時只留 WARNING 以上訊息。
    os.environ.setdefault("LOG_LEVEL", "WARNING")


_bootstrap_backend_import()

from backend.core.templates import MeetingTemplate, get_template, list_templates  # noqa: E402
from backend.core.prompt_templates.section_meeting import TRACKING_TABLE_HEADER  # noqa: E402


# ---------------------------------------------------------------------------
# 由程式碼推導的常數（AST 解析，不 import 服務鏈）
# ---------------------------------------------------------------------------


def read_source(path: Path) -> str:
    """讀取模組原始碼；讀不到時回傳空字串（由呼叫端判定是否致命）。"""
    try:
        return path.read_text(encoding="utf-8")
    except OSError:
        return ""


def iter_module_assignments(source: str) -> Iterator[tuple[str, ast.expr]]:
    """逐一產生模組層級的 `NAME = <expr>` 指派（函式內同名指派不列入）。"""
    if not source:
        return
    tree = ast.parse(source)
    for node in tree.body:
        if not isinstance(node, ast.Assign):
            continue
        for target in node.targets:
            if isinstance(target, ast.Name):
                yield target.id, node.value


def string_assignment(source: str, name: str, prefix: str = "") -> Optional[str]:
    """取出 `name = "<prefix>…"` 形式的字串常數值（模組層級、AST 解析）。"""
    for target_name, value in iter_module_assignments(source):
        if target_name != name:
            continue
        try:
            literal = ast.literal_eval(value)
        except (ValueError, SyntaxError):
            continue
        if isinstance(literal, str) and literal.startswith(prefix):
            return literal
    return None


def first_cjk_phrase(text: str) -> str:
    """取出字串中第一段長度 >= 5 的中文詞（用於推導橫幅關鍵詞）。"""
    match = CJK_PHRASE_PATTERN.search(text)
    return match.group(0) if match else ""


def derive_fallback_markers() -> tuple[str, ...]:
    """由程式碼推導 fallback artifact 的失效標記（不寫死字面值）。

    - 逐字稿 fallback 標題：task_processor 中 `title = "# …"` 指派。
    - 下載檔名標籤：routes.py 的 `record_label = …` 指派。
    - 失敗橫幅關鍵詞：SUMMARY_FAILED_BANNER 內第一段中文詞。
    """
    task_source = read_source(TASK_PROCESSOR_SOURCE_PATH)
    route_source = read_source(RESULT_ROUTE_SOURCE_PATH)

    markers: list[str] = []
    banner = string_assignment(task_source, "SUMMARY_FAILED_BANNER") or ""
    banner_phrase = first_cjk_phrase(banner)
    if banner_phrase:
        markers.append(banner_phrase)
    fallback_title = string_assignment(task_source, "title", prefix="#")
    if fallback_title:
        markers.append(fallback_title.lstrip("# ").strip())
    record_label = string_assignment(route_source, "record_label")
    if record_label:
        markers.append(record_label)

    markers.extend(FALLBACK_TECHNICAL_MARKERS)
    markers.append(FALLBACK_BANNER_EMOJI)
    return tuple(dict.fromkeys(marker for marker in markers if marker))


def tracking_table_columns() -> tuple[str, ...]:
    """由 section_meeting 彙整表表頭常數推導欄位名（待辦／辦理情形欄）。"""
    cells = [cell.strip() for cell in TRACKING_TABLE_HEADER.strip("|").split("|")]
    return tuple(cell for cell in cells if cell)


def general_name_field_pattern() -> tuple[re.Pattern, str]:
    """取得 general 模板的「會議名稱」欄位樣式與其來源說明。

    優先確認該樣式確實出現在 `template.resolve_system_prompt()` 的系統提示詞
    （即 general 模板真正要求的正式欄位），取不到時退回模板樣式本身。
    """
    general = get_template(None)
    for label, pattern in general.required_section_patterns:
        if "會議名稱" in label:
            prompt = general.resolve_system_prompt()
            if pattern.search(prompt):
                return pattern, "「會議名稱：」欄位（general 模板系統提示詞）"
            return pattern, "「會議名稱：」欄位（general 模板 required_section_patterns）"
    raise RuntimeError("general 模板缺少『會議名稱』樣式，無法驗證正式標題")


def template_requires_table(template: MeetingTemplate) -> bool:
    """判斷模板的必備樣式是否以表格為載體（標籤或樣式含「表」）。"""
    return any(
        "表" in label or "表" in pattern.pattern
        for label, pattern in template.required_section_patterns
    )


# ---------------------------------------------------------------------------
# 文字正規化與統計
# ---------------------------------------------------------------------------


def normalize_text(text: str) -> str:
    """正規化全形／半形空白與 Markdown 裝飾，讓 MD 與 DOCX 共用同一套樣式。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u3000", " ").replace("\u00a0", " ")
    text = re.sub(r"^[ \t]*>[ \t]?", "", text, flags=re.MULTILINE)  # 引言標記
    text = re.sub(r"^[ \t]*#{1,6}[ \t]*", "", text, flags=re.MULTILINE)  # 標題井號
    text = re.sub(r"^[ \t]*[-*•][ \t]+", "", text, flags=re.MULTILINE)  # 清單符號
    text = text.replace("**", "").replace("__", "")
    return text


def count_cjk(text: str) -> int:
    """計算中文字數（含延伸 A 與相容表意文字）。"""
    return len(CJK_PATTERN.findall(text))


def strip_placeholders(text: str) -> str:
    """移除「（待確認）」佔位符，用於判斷內容是否實質為空。"""
    return PLACEHOLDER_PATTERN.sub("", text)


def resolve_artifact_template(text: str) -> MeetingTemplate:
    """依紀錄標題解析 artifact 對應的會議模板（無法判定時回退 general）。

    先比對整行標題（Markdown 井號已正規化移除），再以最長 result_title
    內文命中為準；「會議紀錄」為「科務會議紀錄」的子字串，故必須取最長命中。
    """
    normalized = normalize_text(text)
    lines = [line.strip() for line in normalized.splitlines() if line.strip()]
    for template in list_templates():
        title = template.result_title.lstrip("#").strip()
        if not title:
            continue
        if any(line == title or line.startswith(title) for line in lines):
            return template

    best_title = ""
    best_template: Optional[MeetingTemplate] = None
    for template in list_templates():
        title = template.result_title.lstrip("#").strip()
        if title and title in normalized and len(title) > len(best_title):
            best_title, best_template = title, template
    return best_template or get_template(None)


def find_title_line(text: str, title: str) -> Optional[str]:
    """找出等於或起頭於指定標題的行（要求行首命中，避免子字串誤判）。"""
    for line in normalize_text(text).splitlines():
        stripped = line.strip()
        if stripped == title or stripped.startswith(title):
            return stripped
    return None


def window_after(
    text: str, match: re.Match[str], boundary: re.Pattern[str], cap: int = DECISION_WINDOW_CAP
) -> str:
    """取出樣式命中之後的區段文字，遇到章節邊界樣式或長度上限即停止。"""
    tail = text[match.end() : match.end() + cap]
    for boundary_match in boundary.finditer(tail):
        return tail[: boundary_match.start()]
    return tail


def decision_patterns(template: MeetingTemplate) -> tuple[re.Pattern[str], ...]:
    """決議區段樣式：優先取模板必備樣式中標籤含「決議」者，否則用規範樣式。"""
    patterns = tuple(
        pattern for label, pattern in template.required_section_patterns if "決議" in label
    )
    return patterns or (DECISION_HEADING_PATTERN,)


def iter_markdown_tables(lines: Sequence[str]) -> Iterator[list[str]]:
    """逐一產生 Markdown 表格（連續的 | … | 列；規則列已排除）。"""
    table: list[str] = []
    for line in lines:
        if MARKDOWN_TABLE_ROW_PATTERN.match(line):
            if not MARKDOWN_TABLE_RULE_PATTERN.match(line):
                table.append(line)
            continue
        if table:
            yield table
            table = []
    if table:
        yield table


def snippet(text: str) -> str:
    """截斷字串供單行 detail 使用。"""
    compact = re.sub(r"\s+", " ", text.strip())
    return compact[:DETAIL_SNIPPET_LEN]


# ---------------------------------------------------------------------------
# artifact 讀取
# ---------------------------------------------------------------------------


@dataclass
class ArtifactText:
    """單一 artifact 的抽取結果（docx 已含表格列文字）。"""

    kind: str  # "md" | "docx"
    path: Path
    text: str = ""
    error: Optional[str] = None
    paragraph_count: int = 0
    table_count: int = 0
    doc_title: str = ""


def read_markdown_artifact(path: Path) -> ArtifactText:
    """讀取 Markdown artifact。"""
    if not path.is_file():
        return ArtifactText("md", path, error="檔案不存在")
    try:
        return ArtifactText("md", path, path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError) as exc:
        return ArtifactText("md", path, error=f"{type(exc).__name__}: {exc}")


def read_docx_artifact(path: Path) -> ArtifactText:
    """讀取 DOCX artifact：段落與表格列依文件順序串接，表格列轉為 Markdown 列。"""
    if not path.is_file():
        return ArtifactText("docx", path, error="檔案不存在")
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # 環境缺 python-docx：屬環境錯誤，不可放行
        return ArtifactText("docx", path, error=f"python-docx 無法匯入：{exc}")

    try:
        document = Document(str(path))
    except Exception as exc:  # noqa: BLE001 — python-docx 對壞檔拋出多種例外
        return ArtifactText("docx", path, error=f"DOCX 解析失敗（{type(exc).__name__}: {exc}）")

    lines: list[str] = []
    paragraph_count = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph_count += 1
            lines.append(Paragraph(child, document).text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
                lines.append("| " + " | ".join(cells) + " |")

    doc_title = ""
    try:
        doc_title = document.core_properties.title or ""
    except Exception:  # noqa: BLE001 — 讀不到 docProps 不影響其他檢查
        doc_title = ""

    return ArtifactText(
        "docx",
        path,
        "\n".join(lines),
        paragraph_count=paragraph_count,
        table_count=len(document.tables),
        doc_title=doc_title,
    )


# ---------------------------------------------------------------------------
# 檢查結果
# ---------------------------------------------------------------------------


@dataclass
class CheckResult:
    """單一檢查的判定結果。"""

    check_id: str
    status: str  # "PASS" | "FAIL" | "SKIP"
    detail: str

    def render(self) -> str:
        """輸出成單行 `STATUS check-id: detail`。"""
        return f"{self.status} {self.check_id}: {self.detail}"


@dataclass
class ViewFacts:
    """單一 artifact 的正規化事實（各檢查共用，避免重算）。"""

    label: str
    path: Path
    kind: str
    text: str
    template: MeetingTemplate
    error: Optional[str] = None
    paragraph_count: int = 0
    table_count: int = 0
    doc_title: str = ""


def build_facts(artifact: ArtifactText) -> ViewFacts:
    """由讀取結果建立正規化事實；讀取失敗時僅帶錯誤訊息。"""
    if artifact.error is not None:
        return ViewFacts(
            artifact.kind, artifact.path, artifact.kind, "", get_template(None), error=artifact.error
        )
    return ViewFacts(
        label=artifact.kind,
        path=artifact.path,
        kind=artifact.kind,
        text=normalize_text(artifact.text),
        template=resolve_artifact_template(artifact.text),
        paragraph_count=artifact.paragraph_count,
        table_count=artifact.table_count,
        doc_title=artifact.doc_title,
    )


def error_detail(facts: ViewFacts) -> str:
    """組出 artifact 讀取失敗的單行說明。"""
    return f"{facts.path} 無法讀取或解析：{facts.error}"


def combine(check_id: str, findings: Sequence[tuple[str, bool, str]]) -> CheckResult:
    """彙整各 artifact 的判定：任一 FAIL 即 FAIL，否則 PASS。"""
    failed = [(label, detail) for label, passed, detail in findings if not passed]
    if failed:
        return CheckResult(check_id, "FAIL", "；".join(f"{label}：{detail}" for label, detail in failed))
    return CheckResult(check_id, "PASS", "；".join(f"{label}：{detail}" for label, _passed, detail in findings))


# ---------------------------------------------------------------------------
# 各項檢查
# ---------------------------------------------------------------------------


def check_not_fallback(facts_list: Sequence[ViewFacts], markers: Sequence[str]) -> CheckResult:
    """artifact 不得是 fallback 逐字稿（標題／橫幅／技術字樣／docProps 標題）。"""
    findings: list[tuple[str, bool, str]] = []
    for facts in facts_list:
        if facts.error is not None:
            findings.append((facts.label, False, error_detail(facts)))
            continue
        hits = [marker for marker in markers if marker in facts.text]
        title_hit = ""
        if facts.kind == "docx" and facts.doc_title:
            title_hit = next((marker for marker in markers if marker in facts.doc_title), "")
        if hits or title_hit:
            parts: list[str] = []
            if hits:
                parts.append("命中 fallback 標記：" + "、".join(dict.fromkeys(hits)))
            if title_hit:
                parts.append(f"docProps title 即 fallback 標記：{title_hit}")
            findings.append((facts.label, False, "；".join(parts)))
        else:
            findings.append(
                (
                    facts.label,
                    True,
                    "未命中 fallback 標記（標題／⚠️ 橫幅／StableServiceError／Traceback／summary_failed）",
                )
            )
    return combine("not_fallback", findings)


def check_formal_title(
    facts_list: Sequence[ViewFacts], name_pattern: re.Pattern[str], name_source: str
) -> CheckResult:
    """artifact 必須有正式紀錄標題（模板標題行或一般模板的「會議名稱：」欄位）。"""
    findings: list[tuple[str, bool, str]] = []
    for facts in facts_list:
        if facts.error is not None:
            findings.append((facts.label, False, error_detail(facts)))
            continue
        title = facts.template.result_title.lstrip("#").strip()
        title_line = find_title_line(facts.text, title) if title else None
        if title_line:
            findings.append((facts.label, True, f"accepted：標題「{title}」"))
        elif name_pattern.search(facts.text):
            findings.append((facts.label, True, f"accepted：{name_source}"))
        else:
            findings.append((facts.label, False, f"未見正式紀錄標題「{title}」或「會議名稱：」欄位"))
    return combine("formal_title", findings)


def check_required_sections(facts_list: Sequence[ViewFacts]) -> CheckResult:
    """必備章節樣式（含額外欄位）必須全部存在，語意同 `_validate_summary_quality`。"""
    findings: list[tuple[str, bool, str]] = []
    for facts in facts_list:
        if facts.error is not None:
            findings.append((facts.label, False, error_detail(facts)))
            continue
        template = facts.template
        missing = [label for label, pattern in template.required_section_patterns if not pattern.search(facts.text)]
        missing_extra = [
            f"{label}（額外欄位）" for label, pattern in template.extra_field_patterns if not pattern.search(facts.text)
        ]
        total = len(template.required_section_patterns)
        if missing or missing_extra:
            findings.append(
                (
                    facts.label,
                    False,
                    f"{template.id} 模板缺少：" + "、".join(missing + missing_extra),
                )
            )
        else:
            findings.append(
                (facts.label, True, f"{template.id} 模板必備樣式 {total - len(missing)}/{total} 齊備")
            )
    return combine("required_sections", findings)


def check_has_decisions(facts_list: Sequence[ViewFacts]) -> CheckResult:
    """至少一個決議／決議事項區段，且區段內容（扣掉標題本身）有實質長度。"""
    findings: list[tuple[str, bool, str]] = []
    for facts in facts_list:
        if facts.error is not None:
            findings.append((facts.label, False, error_detail(facts)))
            continue
        windows: list[str] = []
        for pattern in decision_patterns(facts.template):
            for match in pattern.finditer(facts.text):
                windows.append(window_after(facts.text, match, facts.template.docx_section_pattern))
        best = max((count_cjk(window) for window in windows), default=0)
        if windows and best >= MIN_DECISION_CJK:
            findings.append(
                (facts.label, True, f"決議區段共 {len(windows)} 處，最長內容 {best} 字（門檻 {MIN_DECISION_CJK}）")
            )
        else:
            findings.append(
                (
                    facts.label,
                    False,
                    f"未見具實質內容的決議區段（偵測 {len(windows)} 處，最大內容 {best} 字，門檻 {MIN_DECISION_CJK}）",
                )
            )
    return combine("has_decisions", findings)


def check_has_action_items(facts_list: Sequence[ViewFacts]) -> CheckResult:
    """至少一個待辦／辦理事項風格的列或表格資料列，且該列有實質內容。"""
    columns = tracking_table_columns()
    findings: list[tuple[str, bool, str]] = []
    for facts in facts_list:
        if facts.error is not None:
            findings.append((facts.label, False, error_detail(facts)))
            continue

        lines = facts.text.splitlines()
        table_hit = ""
        for table in iter_markdown_tables(lines):
            header_hits = [column for column in columns if column in table[0]]
            if len(header_hits) < 2:
                continue
            for row in table[1:]:
                if count_cjk(row) >= MIN_ACTION_TABLE_CJK:
                    table_hit = f"彙整表資料列「{snippet(row)}」（欄位：{'、'.join(header_hits)}）"
                    break
            if table_hit:
                break

        line_hit = ""
        if not table_hit:
            markers: list[tuple[re.Pattern[str], str]] = [
                (pattern, f"模板欄位「{label}」") for label, pattern in facts.template.extra_field_patterns
            ]
            markers.extend(
                (re.compile(re.escape(keyword)), f"待辦關鍵詞「{keyword}」")
                for keyword in SPEC_ACTION_KEYWORDS
            )
            for line in lines:
                if MARKDOWN_TABLE_ROW_PATTERN.match(line):
                    continue  # 表格列由上方彙整表分支處理
                for pattern, label in markers:
                    match = pattern.search(line)
                    if match is None:
                        continue
                    content = strip_placeholders(line[match.end() :])
                    if count_cjk(content) >= MIN_ACTION_CJK:
                        line_hit = f"{label}：{snippet(line)}"
                        break
                if line_hit:
                    break

        if table_hit or line_hit:
            findings.append((facts.label, True, table_hit or line_hit))
        else:
            findings.append((facts.label, False, "未見具內容的待辦／辦理事項列或彙整表資料列"))
    return combine("has_action_items", findings)


def check_min_length(facts_list: Sequence[ViewFacts]) -> CheckResult:
    """正文（含表格文字）至少 MIN_BODY_CJK 個中文字。"""
    findings: list[tuple[str, bool, str]] = []
    for facts in facts_list:
        if facts.error is not None:
            findings.append((facts.label, False, error_detail(facts)))
            continue
        cjk = count_cjk(facts.text)
        if cjk >= MIN_BODY_CJK:
            findings.append((facts.label, True, f"正文 {cjk} 字（門檻 {MIN_BODY_CJK}）"))
        else:
            findings.append((facts.label, False, f"正文僅 {cjk} 字（門檻 {MIN_BODY_CJK}）"))
    return combine("min_length", findings)


def check_no_placeholder(facts_list: Sequence[ViewFacts]) -> CheckResult:
    """不得有 TODO／XXX、不得為「（待確認）」佔位符紀錄、不得全為空章節。"""
    findings: list[tuple[str, bool, str]] = []
    for facts in facts_list:
        if facts.error is not None:
            findings.append((facts.label, False, error_detail(facts)))
            continue

        todo_count = len(TODO_PATTERN.findall(facts.text))
        xxx_count = len(XXX_PATTERN.findall(facts.text))
        meaningful = count_cjk(strip_placeholders(facts.text))

        section_windows: list[str] = []
        for _label, pattern in facts.template.required_section_patterns:
            for match in pattern.finditer(facts.text):
                section_windows.append(window_after(facts.text, match, facts.template.docx_section_pattern))
        has_real_section = any(
            count_cjk(strip_placeholders(window)) >= MIN_SECTION_CJK for window in section_windows
        )

        problems: list[str] = []
        if todo_count:
            problems.append(f"出現 TODO {todo_count} 處")
        if xxx_count:
            problems.append(f"出現 XXX {xxx_count} 處")
        if meaningful < MIN_MEANINGFUL_CJK:
            problems.append(f"扣除（待確認）後實質內容僅 {meaningful} 字（疑似佔位符紀錄）")
        if not has_real_section:
            problems.append("所有必備章節皆無實質內容（空章節）")

        if problems:
            findings.append((facts.label, False, "；".join(problems)))
        else:
            findings.append(
                (facts.label, True, f"無 TODO／XXX，實質內容 {meaningful} 字（非（待確認）佔位符）")
            )
    return combine("no_placeholder", findings)


def check_structure_nonempty(facts_list: Sequence[ViewFacts]) -> CheckResult:
    """DOCX 結構檢查：段落數下限，且模板要求表格時必須真有表格。"""
    docx_facts = [facts for facts in facts_list if facts.kind == "docx"]
    if not docx_facts:
        return CheckResult("structure_nonempty", "SKIP", "未提供 --docx，無法評估 DOCX 結構")

    findings: list[tuple[str, bool, str]] = []
    for facts in docx_facts:
        if facts.error is not None:
            findings.append((facts.label, False, error_detail(facts)))
            continue
        requires_table = template_requires_table(facts.template)
        problems: list[str] = []
        if facts.paragraph_count < MIN_DOCX_PARAGRAPHS:
            problems.append(f"段落僅 {facts.paragraph_count} 段（門檻 {MIN_DOCX_PARAGRAPHS}）")
        if requires_table and facts.table_count == 0:
            problems.append(f"{facts.template.id} 模板要求表格但 DOCX 無表格")
        detail = f"段落 {facts.paragraph_count} 段、表格 {facts.table_count} 個（模板{'要求' if requires_table else '未要求'}表格）"
        findings.append((facts.label, not problems, detail if not problems else "；".join(problems)))
    return combine("structure_nonempty", findings)


def check_cjk_ratio(facts_list: Sequence[ViewFacts]) -> CheckResult:
    """正文須以繁體中文為主，而非英文樣板文字。"""
    findings: list[tuple[str, bool, str]] = []
    for facts in facts_list:
        if facts.error is not None:
            findings.append((facts.label, False, error_detail(facts)))
            continue
        cjk = count_cjk(facts.text)
        latin = len(LATIN_PATTERN.findall(facts.text))
        digits = len(DIGIT_PATTERN.findall(facts.text))
        denominator = cjk + latin + digits
        ratio = cjk / denominator if denominator else 0.0
        detail = f"CJK 占比 {ratio:.3f}（CJK {cjk}／英文 {latin}／數字 {digits}；門檻 {MIN_CJK_RATIO:.2f}）"
        findings.append((facts.label, ratio >= MIN_CJK_RATIO, detail))
    return combine("cjk_ratio", findings)


# ---------------------------------------------------------------------------
# 進入點
# ---------------------------------------------------------------------------


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    """解析命令列參數。"""
    parser = argparse.ArgumentParser(
        description="會議紀錄 artifact（Markdown／DOCX）確定性驗收檢查器",
        epilog="例：uv run python scripts/e2e/check_record_output.py --md out.md --docx out.docx",
    )
    parser.add_argument("--md", type=Path, default=None, help="會議紀錄 Markdown 路徑")
    parser.add_argument("--docx", type=Path, default=None, help="會議紀錄 DOCX 路徑")
    args = parser.parse_args(argv)
    if args.md is None and args.docx is None:
        parser.error("至少需提供 --md 或 --docx 其中一個 artifact 路徑")
    return args


def main(argv: Optional[Sequence[str]] = None) -> int:
    """執行全部檢查並輸出 verdict。"""
    args = parse_args(argv)

    try:
        markers = derive_fallback_markers()
        name_pattern, name_source = general_name_field_pattern()
    except Exception as exc:  # noqa: BLE001 — 推導失敗屬環境錯誤，不得放行
        print(f"ERROR 無法由 backend 程式碼推導檢查依據：{type(exc).__name__}: {exc}", file=sys.stderr)
        return 2

    artifacts: list[ArtifactText] = []
    if args.md is not None:
        artifacts.append(read_markdown_artifact(args.md))
    if args.docx is not None:
        artifacts.append(read_docx_artifact(args.docx))
    facts_list = [build_facts(artifact) for artifact in artifacts]

    results: list[CheckResult] = [
        check_not_fallback(facts_list, markers),
        check_formal_title(facts_list, name_pattern, name_source),
        check_required_sections(facts_list),
        check_has_decisions(facts_list),
        check_has_action_items(facts_list),
        check_min_length(facts_list),
        check_no_placeholder(facts_list),
        check_structure_nonempty(facts_list),
        check_cjk_ratio(facts_list),
    ]

    for result in results:
        print(result.render())

    rejected = any(result.status == "FAIL" for result in results)
    print(f"VERDICT: {'REJECTED' if rejected else 'ACCEPTED'}")
    return 1 if rejected else 0


if __name__ == "__main__":
    sys.exit(main())
