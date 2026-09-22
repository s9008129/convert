# -*- coding: utf-8 -*-
"""T20260827-1127-01 Revision 3 WAVE-01a — E2E runner fail-first acceptance tests。

權威來源：
- plan.md CHANGE_MAP ``CM-01``（九類 fail-first 契約）與 ``IMPLEMENTATION_WAVES`` WAVE-01。
- handoff.md WAVE-02 目標契約（小型可測 helpers、--upload-mode api|browser、
  BROWSER_E2E_READY flush、固定 gate order、fail-closed preflight）。

本模組以 ``importlib`` 載入 ``scripts/e2e/run_owned_e2e.py``（**不修改它**），
用 pure helper、fake httpx/client 與暫存 DOCX/log/資料檔，鎖定 runner「目前缺少」
的驗收行為。每個 red 的失敗原因必須是 runner 缺少批准行為（false-PASS／gate 缺失），
不是 import/fixture/permission 錯誤；WAVE-02 實作後同一 suite 必須轉綠。

安全邊界（WAVE-01 硬性約束）：
- ``subprocess.Popen`` 以 fake 取代——不啟動任何真實 backend child。
- ``httpx`` 以 fake module shim 取代——不出網路、不碰 LM Studio、不碰 9527。
- 不做真實 ASR、不上傳真實音檔；音檔/DOCX/log 全部為 tmp_path 內的合成 fixture。
- ``--audio`` 僅作為 expected-SHA authority 的假檔；run_full_e2e 由 fake client 回應。
"""

import importlib.util
import io
import json
import subprocess  # 僅取用 TimeoutExpired 給 shim，不會啟動任何 process
import sys
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent
RUNNER_PATH = REPO_ROOT / "scripts" / "e2e" / "run_owned_e2e.py"

EXPECTED_REVISION = "expected" + "-rev-0123456789abcdef0123456789abcdef01234567"
STALE_REVISION = "stale" + "-rev-ffffffffffffffffffffffffffffffffffffffff"
MODEL_A_KEY = "qwen3.6-35b-a3b-mlx"
MODEL_B_KEY = "other-model"
INSTANCE_A = "inst-a1"
INSTANCE_B = "inst-b1"
AUDIO_BYTES = b"fake-meeting-audio-bytes-for-e2e-runner-tests"
STORED_FILENAME = "29767e54d905.m4a"
TASK_ID = "task-0001"

FALLBACK_FILENAME_LABEL = "逐字稿(會議紀錄生成失敗)"
FALLBACK_DOCX_TITLE = "逐字稿（會議紀錄生成失敗）"
FALLBACK_ERROR_TEXT = "RuntimeError: 摘要生成失敗：結果為空"
FORMAL_LABEL = "會議紀錄"

UPLOAD_LOG_LINE = (
    f"INFO:     檔案上傳成功: 盤點工具討論.m4a -> {STORED_FILENAME}, 任務ID: {TASK_ID}"
)

METRICS_OK_LINE = (
    "INFO:     本地摘要 pipeline metrics：chunk_count=9, "
    "logical_generations=9, semantic_attempts=9, network_retries=0, "
    "merge_rounds=2, merge_groups_last_round=1, "
    "duration_seconds={'extraction': 100.0, 'merge': 50.0, "
    "'final_and_refine': 20.0, 'total': 170.0}"
)
METRICS_ATTEMPTS_LINE = METRICS_OK_LINE.replace(
    "logical_generations=9, semantic_attempts=9", "logical_generations=1, semantic_attempts=3"
)
METRICS_MERGE_ROUNDS_LINE = METRICS_OK_LINE.replace("merge_rounds=2", "merge_rounds=4")
MAX_TOKENS_1_LINE = (
    "INFO:     使用 LM Studio 生成摘要 (model=qwen3.6-35b-a3b-mlx, instance=inst-a1, "
    "temperature=0.3, max_tokens=1, attempt_type=initial)"
)
HARD_TRUNCATION_LINE = (
    "WARNING:  整併筆記硬截斷：45000 → 3112 tokens（預算 3112）；後段內容未進入最終生成"
)
MERGE_NONCONVERGENCE_LINE = (
    "ERROR:    LM Studio 摘要整併失敗：LOCAL_LLM_MERGE_NOT_CONVERGED（單輪無實質進度）"
)


# ---------------------------------------------------------------------------
# 載入被測 runner（importlib，不修改 scripts/e2e/run_owned_e2e.py）
# ---------------------------------------------------------------------------

def _load_runner_module():
    spec = importlib.util.spec_from_file_location(
        "run_owned_e2e_acceptance_under_test", RUNNER_PATH
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


runner = _load_runner_module()


# ---------------------------------------------------------------------------
# Fake HTTP / process / DOCX fixtures（純模擬，無網路、無真實 process）
# ---------------------------------------------------------------------------

class FakeResponse:
    def __init__(self, status_code=200, json_data=None, content=b"", headers=None, text=None):
        self.status_code = status_code
        self._json = json_data
        self.content = content
        self.headers = headers or {}
        if text is not None:
            self.text = text
        elif json_data is not None:
            self.text = json.dumps(json_data, ensure_ascii=False)
        else:
            self.text = ""

    def json(self):
        if self._json is None:
            raise ValueError("fake response has no JSON body")
        return self._json

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


class FakeChildProc:
    """假 backend child：第一次 poll() 回 None（仍在跑），之後回 0（已結束）。

    runner 的 wait_for_health 只會 poll 一次即拿到 200；finally 的
    terminate_owned_process／child_terminated 觀察到已結束，絕不會真的 killpg。
    """

    def __init__(self, pid=424242):
        self.pid = pid
        self.returncode = 0
        self._polls = 0

    def poll(self):
        self._polls += 1
        return None if self._polls == 1 else 0


class FakeHttpxModule:
    """取代 runner module 內的 httpx；只服務 /api/health 與 /api/v1/models。"""

    HTTPError = httpx.HTTPError

    def __init__(self, *, health_response, models_factory, upload_client):
        self._health_response = health_response
        self._models_factory = models_factory
        self._upload_client = upload_client
        self.models_calls = 0

    def get(self, url, timeout=None, params=None, **kwargs):
        if "/api/health" in url:
            return self._health_response
        if "/api/v1/models" in url:
            self.models_calls += 1
            return self._models_factory(self.models_calls)
        raise AssertionError(f"測試 shim 不預期此 GET：{url}")

    def Client(self, timeout=None, **kwargs):
        return self._upload_client


class FakeUploadClient:
    """模擬 run_full_e2e 的 httpx.Client 與 backend 副作用（upload log + stored bytes）。

    - post /api/upload：記錄呼叫次數（gate-order 斷言用），寫入 backend log 的
      既有「檔案上傳成功 … 任務ID」行，並在 DATA_DIR/uploads 落地 stored bytes。
    - get /api/tasks/<id>：回任務 JSON，並把 metrics/diagnostics 行寫入 backend log
      （模擬 backend 在任務期間已輸出的 structured log）。
    """

    def __init__(
        self,
        *,
        upload_json,
        task_json,
        transcript_response,
        docx_response,
        backend_log_path=None,
        uploads_dir=None,
        stored_filename=STORED_FILENAME,
        stored_bytes=AUDIO_BYTES,
        metrics_lines=(),
    ):
        self._upload_json = upload_json
        self._task_json = task_json
        self._transcript_response = transcript_response
        self._docx_response = docx_response
        self._backend_log_path = backend_log_path
        self._uploads_dir = uploads_dir
        self._stored_filename = stored_filename
        self._stored_bytes = stored_bytes
        self._metrics_lines = list(metrics_lines)
        self.upload_calls = []
        self.task_polls = 0

    def _append_log(self, lines):
        if self._backend_log_path is None:
            return
        self._backend_log_path.parent.mkdir(parents=True, exist_ok=True)
        with open(self._backend_log_path, "a", encoding="utf-8") as fh:
            for line in lines:
                fh.write(line + "\n")

    def post(self, url, files=None, data=None, **kwargs):
        self.upload_calls.append({"url": url, "data": data})
        self._append_log([UPLOAD_LOG_LINE])
        if self._stored_bytes is not None and self._uploads_dir is not None:
            self._uploads_dir.mkdir(parents=True, exist_ok=True)
            (self._uploads_dir / self._stored_filename).write_bytes(self._stored_bytes)
        return FakeResponse(200, self._upload_json)

    def get(self, url, params=None, **kwargs):
        if url.rstrip("/").endswith("/transcript"):
            return self._transcript_response
        if "/result" in url:
            return self._docx_response
        self.task_polls += 1
        self._append_log(self._metrics_lines)
        return FakeResponse(200, self._task_json)


# ---------------------------------------------------------------------------
# DOCX fixtures（python-docx 既有依賴；合成檔，非真實會議內容）
# ---------------------------------------------------------------------------

def _docx_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _fallback_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(f"# {FALLBACK_DOCX_TITLE}", level=1)
    doc.add_paragraph(f"**錯誤訊息**：{FALLBACK_ERROR_TEXT}")
    doc.add_paragraph("（僅逐字稿 fallback 內容）")
    return _docx_bytes(doc)


def _formal_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("盤點工具討論 會議紀錄", level=1)
    doc.add_heading("一、報告事項", level=2)
    doc.add_paragraph("資訊室報告系統上線進度，本週已完成驗收測試。")
    doc.add_heading("二、討論事項", level=2)
    doc.add_paragraph("討論年度採購時程、預算配置與分工。")
    doc.add_paragraph("決議：由資訊室於九月底前完成採購文件並送簽。")
    doc.add_heading("三、主席裁示事項", level=2)
    doc.add_paragraph("主席裁示：照案通過，請各單位配合辦理。")
    return _docx_bytes(doc)


# ---------------------------------------------------------------------------
# 標準 fake flow 組裝
# ---------------------------------------------------------------------------

def _lmstudio_models_response(models_payload):
    return FakeResponse(200, models_payload)


def _one_model_payload():
    return {
        "models": [
            {
                "key": MODEL_A_KEY,
                "type": "llm",
                "loaded_instances": [
                    {"id": INSTANCE_A, "config": {"context_length": 183296}}
                ],
            }
        ]
    }


def _default_task_json(*, summary_failed=False):
    return {
        "task_id": TASK_ID,
        "filename": STORED_FILENAME,
        "original_filename": "盤點工具討論.m4a",
        "file_size": len(AUDIO_BYTES),
        "status": "completed",
        "progress": 100.0,
        "stage": "完成",
        "processing_mode": "local",
        "template_id": "general",
        "summary_failed": summary_failed,
    }


def _upload_json():
    return {
        "task_id": TASK_ID,
        "filename": "盤點工具討論.m4a",
        "file_size": len(AUDIO_BYTES),
        "queue_position": 0,
        "estimated_wait_seconds": 0,
        "message": "檔案上傳成功，已加入處理佇列",
    }


def _transcript_response():
    return FakeResponse(
        200,
        content="會議逐字稿測試內容".encode("utf-8"),
        headers={"content-disposition": 'attachment; filename="20260827120000_逐字稿.txt"'},
    )


def _docx_response(*, label=FORMAL_LABEL, docx_bytes=None):
    return FakeResponse(
        200,
        content=docx_bytes if docx_bytes is not None else _formal_docx_bytes(),
        headers={
            "content-disposition": f'attachment; filename="20260827120000_{label}.docx"'
        },
    )


def _run_runner(
    monkeypatch,
    capsys,
    tmp_path,
    *,
    audio_path,
    data_dir=None,
    health_json=None,
    models_factory=None,
    client=None,
    extra_args=(),
):
    """以 fake stack 執行 runner.main()，回傳可觀察結果（exit code、summary、fakes）。"""
    artifacts_dir = tmp_path / "e2e-artifacts"
    if data_dir is None:
        data_dir = tmp_path / "isolated-data"
    if health_json is None:
        health_json = {"status": "ok", "version": "test", "build_revision": EXPECTED_REVISION}
    if models_factory is None:
        models_factory = lambda _call_index: _lmstudio_models_response(_one_model_payload())
    if client is None:
        client = FakeUploadClient(
            upload_json=_upload_json(),
            task_json=_default_task_json(),
            transcript_response=_transcript_response(),
            docx_response=_docx_response(),
            backend_log_path=artifacts_dir / "backend.log",
            uploads_dir=data_dir / "uploads",
            metrics_lines=[METRICS_OK_LINE],
        )
    fake_proc = FakeChildProc()
    popen_calls = []

    def fake_popen(*args, **kwargs):
        popen_calls.append({"args": args, "kwargs": kwargs})
        return fake_proc

    shim_subprocess = SimpleNamespace(
        Popen=fake_popen,
        TimeoutExpired=subprocess.TimeoutExpired,
        STDOUT=subprocess.STDOUT,
        PIPE=subprocess.PIPE,
        DEVNULL=subprocess.DEVNULL,
    )
    fake_httpx = FakeHttpxModule(
        health_response=FakeResponse(200, health_json),
        models_factory=models_factory,
        upload_client=client,
    )
    monkeypatch.setattr(runner, "subprocess", shim_subprocess)
    monkeypatch.setattr(runner, "httpx", fake_httpx)
    # WAVE-02 clean-worktree gate：以注入點 stub 成 clean（本檔正向控制註解已預期
    # 此注入方式——fake subprocess shim 無 subprocess.run，且測試環境 worktree
    # 是否 dirty 不應干擾 acceptance 語意驗收）。
    monkeypatch.setattr(runner, "git_status_porcelain", lambda: "")
    monkeypatch.setattr(
        sys,
        "argv",
        [
            str(RUNNER_PATH),
            "--artifacts-dir",
            str(artifacts_dir),
            "--expected-revision",
            EXPECTED_REVISION,
            "--audio",
            str(audio_path),
            "--data-dir",
            str(data_dir),
            *extra_args,
        ],
    )

    outcome = {"exit_code": None, "exception": None}
    try:
        outcome["exit_code"] = runner.main()
    except BaseException as exc:  # 含 FileExistsError/PermissionError/SystemExit
        outcome["exception"] = exc

    captured = capsys.readouterr()
    summary = None
    summary_path = artifacts_dir / "run_summary.json"
    if summary_path.exists():
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
    return SimpleNamespace(
        exit_code=outcome["exit_code"],
        exception=outcome["exception"],
        summary=summary,
        artifacts_dir=artifacts_dir,
        popen_calls=popen_calls,
        models_calls=fake_httpx.models_calls,
        client=client,
        output=captured.out + captured.err,
    )


def _standard_flow(
    tmp_path,
    *,
    summary_failed=False,
    docx_label=FORMAL_LABEL,
    docx_bytes=None,
    metrics_lines=(METRICS_OK_LINE,),
    stored_bytes=AUDIO_BYTES,
    health_revision=EXPECTED_REVISION,
    models_factory=None,
    audio_exists=True,
):
    """組出一次 nominal full-mode fake flow 所需的所有元件。"""
    audio_path = tmp_path / "audio.m4a"
    if audio_exists:
        audio_path.write_bytes(AUDIO_BYTES)
    data_dir = tmp_path / "isolated-data"
    artifacts_dir = tmp_path / "e2e-artifacts"
    client = FakeUploadClient(
        upload_json=_upload_json(),
        task_json=_default_task_json(summary_failed=summary_failed),
        transcript_response=_transcript_response(),
        docx_response=_docx_response(label=docx_label, docx_bytes=docx_bytes),
        backend_log_path=artifacts_dir / "backend.log",
        uploads_dir=data_dir / "uploads",
        stored_bytes=stored_bytes,
        metrics_lines=list(metrics_lines),
    )
    return {
        "audio_path": audio_path,
        "data_dir": data_dir,
        "health_json": {
            "status": "ok",
            "version": "test",
            "build_revision": health_revision,
        },
        "models_factory": models_factory
        or (lambda _call: _lmstudio_models_response(_one_model_payload())),
        "client": client,
    }


# ---------------------------------------------------------------------------
# 斷言小工具
# ---------------------------------------------------------------------------

def _failure_reasons(result) -> list:
    if result.summary is None:
        return []
    return list(result.summary.get("failure_reasons", []))


def _reasons_text(result) -> str:
    return "\n".join(_failure_reasons(result)) + "\n" + result.output


def _reasons_match(result, keywords) -> bool:
    return any(kw in reason for reason in _failure_reasons(result) for kw in keywords)


# ---------------------------------------------------------------------------
# CM-01 #1：revision mismatch 必須在任何 upload 之前結束
# ---------------------------------------------------------------------------

def test_revision_mismatch_must_stop_before_any_upload(monkeypatch, capsys, tmp_path):
    """health build_revision != expected 時，upload 呼叫次數必須為 0（fail closed）。

    現行 runner 只把 mismatch append 到 failure_reasons 後繼續 run_full_e2e →
    仍會對 stale backend 執行昂貴上傳（WAVE-01 red 的核心原因）。
    """
    flow = _standard_flow(tmp_path, health_revision=STALE_REVISION)
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )

    assert len(flow["client"].upload_calls) == 0, (
        f"revision mismatch 時 upload 仍被呼叫 {len(flow['client'].upload_calls)} 次——"
        "runner 必須在 upload 前因 mismatch fail closed（CM-01 #1／REQ-CORE-01）"
    )
    assert result.summary is not None and result.summary["verdict"] == "FAIL"
    assert _reasons_match(result, ("revision", "build_revision", "mismatch"))


# ---------------------------------------------------------------------------
# CM-01 #2：completed + summary_failed=true → FAIL
# ---------------------------------------------------------------------------

def test_completed_with_summary_failed_true_must_fail_verdict(monkeypatch, capsys, tmp_path):
    flow = _standard_flow(tmp_path, summary_failed=True)
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )

    assert result.summary is not None, "runner 必須寫出 run_summary.json"
    assert result.summary["verdict"] == "FAIL", (
        f"task=completed + summary_failed=true 時 runner 判 {result.summary['verdict']!r}——"
        "正式會議紀錄驗收必須把 transcript fallback 判為 FAIL（CM-01 #2／RC-3B）"
    )
    assert _reasons_match(result, ("summary_failed", "摘要失敗", "fallback")), (
        "FAIL 原因必須明確指向 summary_failed=true，而非其他無關失敗"
    )


# ---------------------------------------------------------------------------
# CM-01 #3：fallback Content-Disposition / fallback OOXML → FAIL
# ---------------------------------------------------------------------------

def test_fallback_docx_download_label_must_fail(monkeypatch, capsys, tmp_path):
    flow = _standard_flow(
        tmp_path,
        summary_failed=False,  # 隔離文件層偵測：即便旗標缺失/為 False，fallback 文件也必須 FAIL
        docx_label=FALLBACK_FILENAME_LABEL,
        docx_bytes=_fallback_docx_bytes(),
    )
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )

    assert result.summary is not None and result.summary["verdict"] == "FAIL", (
        "Content-Disposition 帶 fallback 下載標籤（逐字稿(會議紀錄生成失敗)）時 runner 仍 PASS——"
        "runner 必須檢查 Content-Disposition 正式 label（CM-01 #3）"
    )
    assert _reasons_match(result, ("fallback", "逐字稿")), (
        "FAIL 原因必須明確指向 fallback 下載標籤"
    )


def test_percent_encoded_formal_download_label_must_pass():
    """Starlette 對非 ASCII 檔名送 RFC 5987 百分比編碼；runner 必須先還原再驗 label。

    實測 attempt-01（2026-09-22）：產品回應
    ``attachment; filename*=utf-8\'\'20260922184230_%E6%9C%83%E8%AD%B0%E7%B4%80%E9%8C%84.docx``
    為**正式**紀錄檔名，舊版 runner 以原始 header 比對字面「會議紀錄」→ 誤判 FAIL。
    """
    disposition = (
        "attachment; filename*=utf-8''"
        "20260922184230_%E6%9C%83%E8%AD%B0%E7%B4%80%E9%8C%84.docx"
    )

    failures = runner.validate_formal_docx_bytes(_formal_docx_bytes(), disposition)

    assert failures == [], f"百分比編碼的正式檔名不得被判失敗：{failures}"


def test_percent_encoded_fallback_download_label_must_still_fail():
    """比例編碼不得成為 fallback 標籤的逃生門。"""
    disposition = (
        "attachment; filename*=utf-8''"
        "20260922184230_%E9%80%90%E5%AD%97%E7%A8%BF%28%E6%9C%83%E8%AD%B0"
        "%E7%B4%80%E9%8C%84%E7%94%9F%E6%88%90%E5%A4%B1%E6%95%97%29.docx"
    )

    failures = runner.validate_formal_docx_bytes(_formal_docx_bytes(), disposition)

    assert any("fallback" in reason for reason in failures), (
        f"百分比編碼的 fallback 檔名必須判 FAIL：{failures}"
    )


def test_section_headings_tolerate_space_after_comma():
    """模板契約為 ``一、\\s*報告事項``；DOCX 章節比對必須容忍頓號後空白。

    實測 attempt-01：模型輸出「一、 報告事項：」為合法契約寫法，舊版字面子字串
    比對讓四條 section 檢查全部誤判為「缺少非空後續內容」。
    """
    from docx import Document

    doc = Document()
    doc.add_heading("0903 科務會議 會議紀錄", level=1)
    doc.add_heading("一、 報告事項：", level=2)
    doc.add_paragraph("組織規程調整由各股配合辦理。")
    doc.add_heading("二、 討論事項：", level=2)
    doc.add_paragraph("文康活動形式討論。")
    doc.add_paragraph("決議：採辦公室下午茶並發禮券。")
    doc.add_heading("三、 主席裁示事項：", level=2)
    doc.add_paragraph("照案通過，請各單位配合。")

    failures = runner.validate_formal_docx_bytes(
        _docx_bytes(doc), "attachment; filename*=utf-8''%E6%9C%83%E8%AD%B0%E7%B4%80%E9%8C%84.docx"
    )

    assert failures == [], f"頓號後空白的章節標題不得被判失敗：{failures}"


def test_fallback_docx_ooxml_content_must_fail(monkeypatch, capsys, tmp_path):
    flow = _standard_flow(
        tmp_path,
        summary_failed=False,  # 隔離 OOXML 內容檢查本身
        docx_label=FORMAL_LABEL,  # header 正式，但 OOXML 內容是 fallback
        docx_bytes=_fallback_docx_bytes(),
    )
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )

    assert result.summary is not None and result.summary["verdict"] == "FAIL", (
        "DOCX OOXML 含 fallback title／empty-summary error 時 runner 仍 PASS——"
        "runner 必須解析 OOXML 文字並拒絕 fallback 文件（CM-01 #3）"
    )
    assert _reasons_match(result, ("fallback", "逐字稿", "OOXML", "摘要生成失敗")), (
        "FAIL 原因必須明確指向 fallback OOXML 內容"
    )


# ---------------------------------------------------------------------------
# CM-01 #4：stored bytes SHA 與 expected audio SHA 不一致 → FAIL
# ---------------------------------------------------------------------------

def test_stored_upload_bytes_sha_mismatch_must_fail(monkeypatch, capsys, tmp_path):
    flow = _standard_flow(tmp_path, stored_bytes=b"tampered-or-wrong-bytes")
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )

    assert result.summary is not None and result.summary["verdict"] == "FAIL", (
        f"stored upload bytes={flow['client']._stored_bytes!r} 與 source audio 不同仍 PASS——"
        "runner 必須對 DATA_DIR/uploads/<stored> 計算 SHA-256 並與 expected 比對（CM-01 #4）"
    )
    assert _reasons_match(result, ("sha", "SHA", "hash", "雜湊", "雜湊值")), (
        "FAIL 原因必須明確指向 stored bytes SHA mismatch"
    )


def test_sha_match_positive_control_is_not_vetoed(monkeypatch, capsys, tmp_path):
    """正向控制：stored bytes 與 source audio 一致時，SHA gate 不得產生任何失敗原因。"""
    flow = _standard_flow(tmp_path, stored_bytes=AUDIO_BYTES)
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )

    sha_reasons = [
        reason
        for reason in _failure_reasons(result)
        if any(kw in reason for kw in ("sha", "SHA", "hash", "雜湊"))
    ]
    assert not sha_reasons, f"SHA 一致時不得出現 SHA gate 失敗原因：{sha_reasons}"


# ---------------------------------------------------------------------------
# CM-01 #5：model/instance 非唯一，或 run 前後 snapshot 改變 → FAIL
# ---------------------------------------------------------------------------

def _assert_unique_model_gate(result, case_name):
    assert result.summary is not None and result.summary["verdict"] == "FAIL", (
        f"loaded LLM inventory 非唯一（{case_name}）時 runner 仍 PASS——"
        "runner 必須要求恰好一個 model 且恰好一個 instance（CM-01 #5）"
    )
    assert _reasons_match(result, ("model", "instance", "inventory")), (
        "FAIL 原因必須明確指向 model/instance 唯一性"
    )
    assert len(result.client.upload_calls) == 0, (
        "model gate 失敗時不得進入 upload（gate ordering：unique model → upload）"
    )


def test_two_loaded_llm_models_must_fail_before_upload(monkeypatch, capsys, tmp_path):
    models_payload = {
        "models": [
            {
                "key": MODEL_A_KEY,
                "type": "llm",
                "loaded_instances": [{"id": INSTANCE_A, "config": {"context_length": 183296}}],
            },
            {
                "key": MODEL_B_KEY,
                "type": "llm",
                "loaded_instances": [{"id": INSTANCE_B, "config": {"context_length": 8192}}],
            },
        ]
    }
    flow = _standard_flow(tmp_path, models_factory=lambda _c: _lmstudio_models_response(models_payload))
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )
    _assert_unique_model_gate(result, "two_models")


def test_two_loaded_instances_of_one_model_must_fail_before_upload(monkeypatch, capsys, tmp_path):
    models_payload = {
        "models": [
            {
                "key": MODEL_A_KEY,
                "type": "llm",
                "loaded_instances": [
                    {"id": INSTANCE_A, "config": {"context_length": 183296}},
                    {"id": INSTANCE_B, "config": {"context_length": 8192}},
                ],
            }
        ]
    }
    flow = _standard_flow(tmp_path, models_factory=lambda _c: _lmstudio_models_response(models_payload))
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )
    _assert_unique_model_gate(result, "two_instances")


def test_zero_loaded_llm_model_must_fail_before_upload(monkeypatch, capsys, tmp_path):
    flow = _standard_flow(tmp_path, models_factory=lambda _c: _lmstudio_models_response({"models": []}))
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )
    _assert_unique_model_gate(result, "zero_models")


def test_model_changed_between_start_and_end_snapshots_must_fail(monkeypatch, capsys, tmp_path):
    """start snapshot 是唯一 model A；end snapshot 變成 model B → 必須 FAIL。

    現行 runner 只 capture 一次 snapshot 且不比較（capability absent + 行為缺失）。
    """
    start_payload = {
        "models": [
            {
                "key": MODEL_A_KEY,
                "type": "llm",
                "loaded_instances": [{"id": INSTANCE_A, "config": {"context_length": 183296}}],
            }
        ]
    }
    end_payload = {
        "models": [
            {
                "key": MODEL_B_KEY,
                "type": "llm",
                "loaded_instances": [{"id": "inst-b1", "config": {"context_length": 8192}}],
            }
        ]
    }

    def models_factory(call_index):
        return _lmstudio_models_response(start_payload if call_index == 1 else end_payload)

    flow = _standard_flow(tmp_path, models_factory=models_factory)
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )

    assert result.summary is not None and result.summary["verdict"] == "FAIL", (
        f"run 前後 model snapshot 改變（{MODEL_A_KEY} → {MODEL_B_KEY}）仍 PASS——"
        "runner 必須 capture 並比較 ending model snapshot（CM-01 #5／CM-02 metrics+model snapshot）"
    )
    assert _reasons_match(result, ("model", "instance", "snapshot")), (
        "FAIL 原因必須明確指向 model snapshot 前後不一致"
    )


# ---------------------------------------------------------------------------
# CM-01 #6：invalid/unwritable DATA_DIR → backend 不得啟動且明確 fail
# ---------------------------------------------------------------------------

def _looks_like_explicit_data_dir_gate_failure(result) -> bool:
    combined = _reasons_text(result)
    mentions_data_dir = "DATA_DIR" in combined or "data_dir" in combined or "data dir" in combined.lower()
    mentions_gate_semantics = any(
        kw in combined
        for kw in (
            "不可寫", "無法寫入", "可寫", "writable", "寫入",
            "無效", "invalid", "不是目錄", "必須為目錄", "必須是目錄",
            "無法建立", "不可建立", "must be a directory", "not writable",
        )
    )
    return mentions_data_dir and mentions_gate_semantics


@pytest.mark.parametrize("case", ["data-dir-is-regular-file", "data-dir-parent-readonly"])
def test_invalid_or_unwritable_data_dir_fails_closed(monkeypatch, capsys, tmp_path, case):
    audio_path = tmp_path / "audio.m4a"
    audio_path.write_bytes(AUDIO_BYTES)

    read_only_parent = None
    if case == "data-dir-is-regular-file":
        bad_data_dir = tmp_path / "not-a-dir"
        bad_data_dir.write_bytes(b"this is a file, not a directory")
    else:
        read_only_parent = tmp_path / "read-only"
        read_only_parent.mkdir()
        bad_data_dir = read_only_parent / "data"
    try:
        if read_only_parent is not None:
            read_only_parent.chmod(0o555)
        result = _run_runner(
            monkeypatch, capsys, tmp_path,
            audio_path=audio_path, data_dir=bad_data_dir,
        )
    finally:
        if read_only_parent is not None:
            read_only_parent.chmod(0o755)

    assert len(result.popen_calls) == 0, (
        f"invalid/unwritable DATA_DIR（case={case}）時 backend 仍被啟動——"
        "DATA_DIR preflight 必須在啟動 child 之前 fail closed（CM-01 #6）"
    )
    assert _looks_like_explicit_data_dir_gate_failure(result), (
        f"invalid/unwritable DATA_DIR（case={case}）必須輸出明確的 DATA_DIR gate failure"
        "（例如 exit 1 + 提及 DATA_DIR 不可寫/無效的訊息或 explicit RuntimeError）；"
        f"實際：exception={result.exception!r}, summary={result.summary!r}, output={result.output!r}"
    )
    assert result.exit_code == 1 or result.exception is not None


# ---------------------------------------------------------------------------
# CM-01（gate order）：audio 缺檔 → 不得啟動 backend
# ---------------------------------------------------------------------------

def test_missing_audio_fails_closed_before_backend_start(monkeypatch, capsys, tmp_path):
    flow = _standard_flow(tmp_path, audio_exists=False)
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
    )

    assert len(result.popen_calls) == 0, (
        "--audio 不存在時 backend 仍被啟動（Popen 呼叫 "
        f"{len(result.popen_calls)} 次）——audio 存在性屬啟動前 hard preflight（CM-01 #6／固定 gate order）"
    )
    assert _reasons_match(result, ("audio", "音檔")), (
        "必須輸出明確的 audio 缺失 failure，而非深層 traceback"
    )


# ---------------------------------------------------------------------------
# CM-01 #7：Browser upload-log mapping 必須唯一（0 筆 / >1 筆皆 FAIL）
# ---------------------------------------------------------------------------

def _callable_names(module):
    return [name for name in dir(module) if callable(getattr(module, name))]


def _find_browser_upload_mapping_helper(module):
    """尋找「backend log → (stored filename, task_id)」唯一映射解析 helper。"""
    for name in _callable_names(module):
        lowered = name.lower()
        if "upload" in lowered and ("log" in lowered or "mapping" in lowered or "browser" in lowered):
            return getattr(module, name), name
    return None, None


def _call_mapping_helper(helper, log_text):
    try:
        return helper(log_text)
    except TypeError:
        return helper(log_text.splitlines())


def _coerce_mapping_pair(result):
    if isinstance(result, (tuple, list)) and len(result) == 2:
        return str(result[0]), str(result[1])
    if isinstance(result, dict):
        stored = next((v for k, v in result.items() if "filename" in k.lower() or "stored" in k.lower()), None)
        task = next((v for k, v in result.items() if "task" in k.lower()), None)
        if stored is not None and task is not None:
            return str(stored), str(task)
    return None


def _mapping_helper_rejects(helper, log_text):
    try:
        result = _call_mapping_helper(helper, log_text)
    except Exception:
        return True
    if result is None or result is False:
        return True
    return _coerce_mapping_pair(result) is None


def test_browser_upload_log_mapping_requires_unique_match():
    single = (
        f"2026-08-27 12:00:00 | INFO | 檔案上傳成功: 盤點工具討論.m4a -> {STORED_FILENAME}, "
        f"任務ID: {TASK_ID}\n"
    )
    zero = "2026-08-27 12:00:01 | INFO | 與上傳無關的訊息\n"
    double = single + single.replace(TASK_ID, "task-0002")

    helper, name = _find_browser_upload_mapping_helper(runner)
    assert helper is not None, (
        "WAVE-02 CM-01 #7：runner 必須提供『從 isolated backend.log bounded 解析唯一 "
        "stored filename + task_id』的 helper（0 筆逾時、>1 筆皆 FAIL）；"
        f"現行 runner 沒有任何 upload-log mapping 解析能力（callables：{_callable_names(runner)}）"
    )

    # 行為契約（WAVE-02 helper 落地後生效）：唯一一筆 → 成功；0 / >1 筆 → 判失敗
    result = _call_mapping_helper(helper, single)
    pair = _coerce_mapping_pair(result)
    assert pair == (STORED_FILENAME, TASK_ID), (
        f"唯一 upload mapping 應解析出 ({STORED_FILENAME!r}, {TASK_ID!r})，實際：{result!r}"
    )
    assert _mapping_helper_rejects(helper, zero), "0 筆 upload mapping 必須判失敗/逾時"
    assert _mapping_helper_rejects(helper, double), ">1 筆 upload mapping 必須判 FAIL"


def test_upload_log_mapping_strips_ansi_color_codes():
    """CM-01 #7 補強：backend loguru colorize=True 寫檔，log 行含 ANSI 色碼。

    attempt-01 事故：task_id 尾端黏上 \\x1b[0m → httpx InvalidURL →
    runner 提前終止（ASR 子程序連帶 SIGTERM）。解析必須先剝除 ANSI。
    """
    ansi_line = (
        "2026-08-27 21:40:08 | INFO     | backend.api.routes:upload_file:187 - "
        "\x1b[1m檔案上傳成功\x1b[0m: 盤點工具討論.m4a -> "
        f"{STORED_FILENAME}\x1b[0m, 任務ID: {TASK_ID}\x1b[0m\n"
    )
    matches = runner.parse_upload_success_lines(ansi_line)
    assert len(matches) == 1, f"ANSI 彩色 log 行應解析出唯一 mapping，實際：{matches!r}"
    parsed = matches[0]
    for key in ("original_filename", "stored_filename", "task_id"):
        assert "\x1b" not in parsed[key], f"{key} 不得殘留 ANSI 序列：{parsed[key]!r}"
    assert parsed["stored_filename"] == STORED_FILENAME
    assert parsed["task_id"] == TASK_ID, (
        f"task_id 應為乾淨 {TASK_ID!r}，實際 {parsed['task_id']!r}"
        "（ANSI 色碼殘留會使後續 URL 建構丟 InvalidURL）"
    )


# ---------------------------------------------------------------------------
# CM-01 #8：structured metrics 驗收 gate
# ---------------------------------------------------------------------------

def _assert_metrics_veto(result, case):
    assert result.summary is not None and result.summary["verdict"] == "FAIL", (
        f"metrics case={case!r} 時 runner 仍 PASS——runner 必須解析 structured metrics 並"
        "拒絕缺失/超界/截斷/non-convergence 成功假象（CM-01 #8）"
    )
    metric_keywords = (
        "metric", "metrics", "logical_generations", "semantic_attempts",
        "merge_rounds", "max_tokens", "截斷", "truncat", "converge", "CONVERGED",
    )
    assert _reasons_match(result, metric_keywords), (
        f"FAIL 原因必須明確指向 metrics 違規（case={case!r}）；實際 reasons：{_failure_reasons(result)}"
    )


@pytest.mark.parametrize(
    "case, metrics_lines",
    [
        ("missing-metrics", []),
        ("semantic-attempts-above-bound", [METRICS_ATTEMPTS_LINE]),
        ("merge-rounds-exceed-3", [METRICS_MERGE_ROUNDS_LINE]),
        ("max-tokens-1", [METRICS_OK_LINE, MAX_TOKENS_1_LINE]),
        ("hard-truncation", [METRICS_OK_LINE, HARD_TRUNCATION_LINE]),
        ("merge-nonconvergence-while-completed", [METRICS_OK_LINE, MERGE_NONCONVERGENCE_LINE]),
    ],
)
def test_invalid_metrics_must_fail(monkeypatch, capsys, tmp_path, case, metrics_lines):
    flow = _standard_flow(tmp_path, metrics_lines=metrics_lines)
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )
    _assert_metrics_veto(result, case)


def test_valid_bounded_metrics_do_not_veto(monkeypatch, capsys, tmp_path):
    """正向控制：有界且收斂的 metrics 不得被 metrics gate 否決。"""
    flow = _standard_flow(tmp_path, metrics_lines=[METRICS_OK_LINE])
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )

    metric_keywords = (
        "metric", "metrics", "logical_generations", "semantic_attempts",
        "merge_rounds", "max_tokens", "截斷", "truncat", "converge", "CONVERGED",
    )
    offending = [r for r in _failure_reasons(result) if any(kw in r for kw in metric_keywords)]
    assert not offending, f"有效 bounded metrics 不得被 veto：{offending}"
    assert result.summary is not None and result.summary["verdict"] == "PASS", (
        f"nominal fake flow（含有效 metrics）應 PASS；實際 reasons：{_failure_reasons(result)}"
    )


# ---------------------------------------------------------------------------
# CM-01 #9：gate helpers 存在且在 backend start／upload 之前執行
# ---------------------------------------------------------------------------

def test_preflight_gate_helpers_exist_and_run_before_start_and_upload(monkeypatch, capsys, tmp_path):
    """批准契約：revision/clean/data/model 等 gates 必須是可測試的 discrete helpers，
    且在啟動 backend／任何 upload 之前執行（行為面另由本檔其他測試的
    popen/upload call count 斷言鎖定）。"""
    names = _callable_names(runner)
    validate_verbs = ("validate", "check", "ensure", "preflight", "gate", "assert")

    data_gates = [n for n in names if "data" in n.lower() and any(v in n.lower() for v in validate_verbs)]
    clean_gates = [n for n in names if any(v in n.lower() for v in ("clean", "dirty", "worktree"))]
    model_gates = [
        n for n in names
        if "model" in n.lower() and any(v in n.lower() for v in (*validate_verbs, "unique", "consistency", "snapshot_match"))
    ]
    metrics_helpers = [n for n in names if "metric" in n.lower()]
    sha_helpers = [n for n in names if "sha" in n.lower() or "hash" in n.lower()]

    missing = []
    if not data_gates:
        missing.append("DATA_DIR 可寫/有效 preflight helper（如 validate_data_dir）")
    if not clean_gates:
        missing.append("clean worktree / revision gate helper")
    if not model_gates:
        missing.append("model/instance 唯一性 gate helper")
    if not metrics_helpers:
        missing.append("structured metrics validator/parser")
    if not sha_helpers:
        missing.append("SHA-256 helper（source audio vs stored bytes 比對）")
    assert not missing, (
        "WAVE-02 gate 契約缺少可測試 helpers（gate 必須存在且在 backend start/upload 前執行）："
        + "; ".join(missing)
        + f"；現行 runner callables：{names}"
    )


# ---------------------------------------------------------------------------
# 正向控制：nominal fake flow 不得被任何 gate 誤殺
# ---------------------------------------------------------------------------

def test_nominal_formal_run_passes_all_gates(monkeypatch, capsys, tmp_path):
    """正向控制：全部條件合規（formal docx、SHA 一致、唯一 model、有效 metrics）
    時必須 PASS——防止 WAVE-02 的 gates 反向過度阻擋。

    注意：若 WAVE-02 加入真實 `git status` 型 repo-clean gate，請把它做成
    可注入 pure helper 並在本測試 monkeypatch，避免測試環境 dirty worktree 干擾。
    """
    flow = _standard_flow(tmp_path, stored_bytes=AUDIO_BYTES, metrics_lines=[METRICS_OK_LINE])
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        models_factory=flow["models_factory"],
        client=flow["client"],
    )

    assert result.summary is not None, "nominal flow 必須寫出 run_summary.json"
    assert result.summary["verdict"] == "PASS", (
        f"nominal fake flow 應 PASS，實際 reasons：{_failure_reasons(result)}"
    )


# ---------------------------------------------------------------------------
# v4.8.1：驗收場景必須等於使用場景（--template 須送出且須相符）
# ---------------------------------------------------------------------------

def _template_flow(tmp_path, *, task_template_id):
    """組出一次帶自訂 template 的 fake flow（client 的 task 回傳指定模板）。"""
    flow = _standard_flow(tmp_path)
    client = FakeUploadClient(
        upload_json=_upload_json(),
        task_json={**_default_task_json(), "template_id": task_template_id},
        transcript_response=_transcript_response(),
        # W8 後 DOCX 正式性檢查為模板感知：section_meeting 場次的 fixture 必須是
        # 該模板的真實形狀（4 個模板章節），否則 nominal flow 會因 fixture 與
        # 場景不一致而 FAIL（此非產品缺陷）。
        docx_response=_docx_response(docx_bytes=_section_meeting_docx_bytes()),
        backend_log_path=flow["data_dir"].parent / "e2e-artifacts" / "backend.log",
        uploads_dir=flow["data_dir"] / "uploads",
        metrics_lines=[METRICS_OK_LINE],
    )
    flow["client"] = client
    return flow


def test_template_argument_is_posted_and_verified(monkeypatch, capsys, tmp_path):
    """`--template section_meeting` 時：表單須帶 meeting_template，且
    task_final.template_id 相符才 PASS（驗收場不得誤用預設 general）。"""
    flow = _template_flow(tmp_path, task_template_id="section_meeting")
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        client=flow["client"],
        extra_args=("--template", "section_meeting"),
    )

    assert flow["client"].upload_calls, "必須真的上傳音檔"
    posted = flow["client"].upload_calls[0]["data"] or {}
    assert posted.get("meeting_template") == "section_meeting", (
        f"上傳表單未帶 meeting_template，實際 data={posted!r}——"
        "驗收場若不指定模板，產物品質與使用者真實路徑（section_meeting）不同"
    )
    assert result.summary is not None, "runner 必須寫出 run_summary.json"
    assert result.summary.get("meeting_template") == "section_meeting", (
        "run_summary 必須記錄本次驗收使用的模板，供事後判定可比性"
    )
    assert result.summary["verdict"] == "PASS", (
        f"模板相符的 nominal flow 應 PASS，實際 reasons：{_failure_reasons(result)}"
    )


def test_template_mismatch_must_fail_verdict(monkeypatch, capsys, tmp_path):
    """要求 section_meeting 但任務實際用 general → 必須 FAIL（場景不一致）。"""
    flow = _template_flow(tmp_path, task_template_id="general")
    result = _run_runner(
        monkeypatch, capsys, tmp_path,
        audio_path=flow["audio_path"],
        data_dir=flow["data_dir"],
        health_json=flow["health_json"],
        client=flow["client"],
        extra_args=("--template", "section_meeting"),
    )

    assert result.summary is not None, "runner 必須寫出 run_summary.json"
    assert result.summary["verdict"] == "FAIL", (
        "模板未生效時仍判 PASS，等於用錯誤場景的產物背書品質"
    )
    assert _reasons_match(result, ("模板未生效", "meeting_template")), (
        f"FAIL 原因必須指向模板不一致，實際：{_failure_reasons(result)}"
    )


# ---------------------------------------------------------------------------
# W8（plan rev6）：DOCX 正式性檢查模板感知（section_meeting 不再被 general 章節誤殺）
# ---------------------------------------------------------------------------

FORMAL_DISPOSITION = (
    "attachment; filename*=utf-8''20260922184230_%E6%9C%83%E8%AD%B0%E7%B4%80%E9%8C%84.docx"
)


def _section_meeting_docx_bytes() -> bytes:
    """合法 section_meeting 紀錄：具 4 個模板章節、無 general 章節。"""
    from docx import Document

    doc = Document()
    doc.add_heading("0903 科務會議紀錄", level=1)
    doc.add_paragraph("時間：中華民國115年9月3日")
    doc.add_paragraph("主持人：科長")
    doc.add_paragraph("決議事項：")
    doc.add_paragraph("案由及承辦單位：組織異動相關系統權限調整（科長，00:05:36）")
    doc.add_paragraph("一、科長轉知局務會議工作報告及相關注意事項：")
    doc.add_paragraph("1. 土地稅科分拆為地價稅科與土地增值稅科。（科長，00:05:36）")
    doc.add_paragraph("二、科長指示及提醒事項：")
    doc.add_paragraph("1. 全體同仁設定郵件為文字模式。（科長，00:30:33）")
    doc.add_paragraph("散會：下午五時三十分。")
    return _docx_bytes(doc)


def _general_less_docx_bytes() -> bytes:
    """含正式標題、但完全沒有 general 4 章節的合成 DOCX。"""
    from docx import Document

    doc = Document()
    doc.add_heading("0903 科務會議紀錄", level=1)
    doc.add_paragraph("時間：中華民國115年9月3日")
    doc.add_paragraph("案由及承辦單位：組織異動調整")
    doc.add_paragraph("散會：下午五時三十分。")
    return _docx_bytes(doc)


GENERAL_SECTION_FAILURE_TEMPLATE = "DOCX OOXML 的 general 模板 section「{}」缺少非空後續內容"


def test_section_meeting_docx_passes_template_sections_without_general_sections():
    """section_meeting 完整紀錄（無 general 章節）→ 模板感知檢查必須 PASS。

    attempt-03 基線：runner 寫死 general 章節，導致完整科務會議紀錄 DOCX 永遠
    被判 FAIL，`--template section_meeting` 場次不可能 PASS（plan W8）。
    """
    failures = runner.validate_formal_docx_bytes(
        _section_meeting_docx_bytes(), FORMAL_DISPOSITION, template_id="section_meeting"
    )

    assert failures == [], f"section_meeting 4 章節齊備時不得因缺 general 章節被判失敗：{failures}"


def test_general_and_unspecified_template_keep_exact_previous_behaviour(capsys):
    """general／None／未列表模板：失敗字串與修正前逐字元相同（不得被 W8 改壞）。"""
    docx_bytes = _general_less_docx_bytes()
    expected = [
        GENERAL_SECTION_FAILURE_TEMPLATE.format(section)
        for section in runner.GENERAL_REQUIRED_SECTIONS
    ]

    failures_none = runner.validate_formal_docx_bytes(docx_bytes, FORMAL_DISPOSITION)
    failures_general = runner.validate_formal_docx_bytes(
        docx_bytes, FORMAL_DISPOSITION, template_id="general"
    )
    failures_unlisted = runner.validate_formal_docx_bytes(
        docx_bytes, FORMAL_DISPOSITION, template_id="procurement_evaluation"
    )
    captured = capsys.readouterr()

    assert failures_none == expected, f"未指定模板行為必須不變：{failures_none}"
    assert failures_general == expected, f"general 行為必須不變：{failures_general}"
    assert failures_unlisted == expected, (
        f"未列表模板沿用 general 契約（現行行為不變）：{failures_unlisted}"
    )
    assert "[WARN]" in captured.err and "procurement_evaluation" in captured.err, (
        "未列表模板必須另記一行 log 說明沿用 general 契約（已知限制）"
    )
