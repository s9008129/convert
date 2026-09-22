# -*- coding: utf-8 -*-
"""T20260922-2037-02 P4-D（W3）：runner 品質閘門／provider 拆分／mode_used 斷言測試。

權威來源：
- handoff.md §4.4（W3 runner 契約）：``--quality-mode {off,observe,required}``（預設
  off）；``off``＝``run_summary.json`` 不得出現新鍵且 verdict 與現行相同；
  ``observe``＝品質有值、verdict 不因品質值改變；``required``＝低品質 fixture 必
  FAIL 且 ``failure_reasons`` 逐條指名（無 ``--template``＝``parser.error``）；
  門檻一律沿用既有值（不得新增／放寬）；``mode_used`` 事後斷言不符 ⇒ FAIL；
  required 清單拆「共通必要／provider 相依」。
- handoff.md §11 C2：required 計數為 smoke 4／full 14／full+template 15；
  ``run_summary.checks`` 16 鍵＝15 required＋``model_snapshot_captured``。
- handoff.md §11 N3：``off`` 的 byte 級不變以**函式級 golden** 驗證（本檔以
  golden checks 全集＋summary 頂層鍵集合＋verdict 鎖定），不得只靠兩次 live。

安全邊界（沿用 ``tests/test_owned_e2e_acceptance.py`` 慣例）：
- ``subprocess.Popen``／``httpx`` 以 fake shim 取代——不啟動 backend、不出網路、
  不碰 LM Studio、不碰 9527；品質儀器以 monkeypatch fake 取代（只寫 ``--out``
  JSON）。唯一真 subprocess 測試為 ``run_quality_instrument`` 的 fail-soft 單元測試。
- 品質 fixture 為合成數值；錄音／DOCX／log 皆 tmp_path 內合成物。
"""

import hashlib
import importlib.util
import io
import json
import subprocess  # 僅供 shim 取用 TimeoutExpired；不會啟動真實 backend
import sys
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent
RUNNER_PATH = REPO_ROOT / "scripts" / "e2e" / "run_owned_e2e.py"

EXPECTED_REVISION = "expected" + "-rev-0123456789abcdef0123456789abcdef01234567"
AUDIO_BYTES = b"fake-meeting-audio-bytes-for-p4d-quality-gate-tests"
STORED_FILENAME = "29767e54d905.m4a"
TASK_ID = "task-p4d-0001"
MODEL_KEY = "qwen3.6-35b-a3b-mlx"
INSTANCE_ID = "inst-p4d-1"
FORMAL_LABEL = "會議紀錄"
RECORD_MARKDOWN_NAME = "0903-科務會議_" + TASK_ID + ".md"

UPLOAD_LOG_LINE = (
    "INFO:     檔案上傳成功: 0903-科務會議.m4a -> " + STORED_FILENAME + ", 任務ID: " + TASK_ID
)
METRICS_OK_LINE = (
    "INFO:     本地摘要 pipeline metrics：chunk_count=9, "
    "logical_generations=9, semantic_attempts=9, network_retries=0, "
    "merge_rounds=2, merge_groups_last_round=1, "
    "duration_seconds={'extraction': 100.0, 'merge': 50.0, "
    "'final_and_refine': 20.0, 'total': 170.0}"
)
TAG_METRIC_VERSION = "tag_traceability-1.1.0"

# P4-D golden（off 模式＝與現行完全一致；C2：full 無 template 為 15 個 checks 鍵）
GOLDEN_SUMMARY_KEYS = {
    "mode",
    "upload_mode",
    "meeting_template",
    "started_at",
    "finished_at",
    "port",
    "child_pid",
    "expected_build_revision",
    "actual_build_revision",
    "artifacts_dir",
    "runtime_dir",
    "backend_log_path",
    "data_dir",
    "source_audio_sha256",
    "checks",
    "failure_reasons",
    "verdict",
}
GOLDEN_FULL_CHECKS = {
    "backend_started": True,
    "health_ok": True,
    "build_revision_match": True,
    "model_snapshot_captured": True,
    "model_inventory_unique": True,
    "upload_ok": True,
    "stored_upload_sha_match": True,
    "task_completed": True,
    "task_summary_failed_false": True,
    "transcript_downloaded": True,
    "docx_downloaded": True,
    "formal_docx_valid": True,
    "metrics_valid": True,
    "model_snapshot_consistent": True,
    "child_terminated": True,
}


def _load_runner_module():
    spec = importlib.util.spec_from_file_location("run_owned_e2e_p4d_under_test", RUNNER_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


runner = _load_runner_module()


# ---------------------------------------------------------------------------
# Fake HTTP / process fixtures（純模擬；無網路、無真實 process）
# ---------------------------------------------------------------------------

class FakeResponse:
    def __init__(self, status_code=200, json_data=None, content=b"", headers=None, text=None):
        self.status_code = status_code
        self._json = json_data
        self.content = content
        self.headers = headers or {}
        if text is not None:
            self.text = text
        elif json_data is not None:
            self.text = json.dumps(json_data, ensure_ascii=False)
        else:
            self.text = ""

    def json(self):
        if self._json is None:
            raise ValueError("fake response has no JSON body")
        return self._json

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError("HTTP " + str(self.status_code))


class FakeChildProc:
    """假 backend child：第一次 poll() 回 None（仍在跑），之後回 0（已結束）。"""

    def __init__(self, pid=515151):
        self.pid = pid
        self.returncode = 0
        self._polls = 0

    def poll(self):
        self._polls += 1
        return None if self._polls == 1 else 0


class FakeHttpxModule:
    """取代 runner module 內的 httpx；服務 /api/health、/api/v1/models、/api/config。"""

    HTTPError = httpx.HTTPError

    def __init__(self, *, health_response, models_factory, config_response, upload_client):
        self._health_response = health_response
        self._models_factory = models_factory
        self._config_response = config_response
        self._upload_client = upload_client
        self.models_calls = 0
        self.config_calls = 0

    def get(self, url, timeout=None, params=None, **kwargs):
        if "/api/health" in url:
            return self._health_response
        if "/api/v1/models" in url:
            self.models_calls += 1
            return self._models_factory(self.models_calls)
        if "/api/config" in url:
            self.config_calls += 1
            if isinstance(self._config_response, Exception):
                raise self._config_response
            return self._config_response
        raise AssertionError("測試 shim 不預期此 GET：" + url)

    def Client(self, timeout=None, **kwargs):
        return self._upload_client


class FakeUploadClient:
    """模擬 run_full_e2e 的 httpx.Client 與 backend 副作用（upload log／stored
    bytes／逐字稿／DOCX／結果 Markdown）。"""

    def __init__(
        self,
        *,
        upload_json,
        task_json,
        transcript_response,
        docx_response,
        backend_log_path=None,
        uploads_dir=None,
        outputs_dir=None,
        stored_filename=STORED_FILENAME,
        stored_bytes=AUDIO_BYTES,
        metrics_lines=(),
    ):
        self._upload_json = upload_json
        self._task_json = task_json
        self._transcript_response = transcript_response
        self._docx_response = docx_response
        self._backend_log_path = backend_log_path
        self._uploads_dir = uploads_dir
        self._outputs_dir = outputs_dir
        self._stored_filename = stored_filename
        self._stored_bytes = stored_bytes
        self._metrics_lines = list(metrics_lines)
        self.upload_calls = []
        self.task_polls = 0

    def _append_log(self, lines):
        if self._backend_log_path is None:
            return
        self._backend_log_path.parent.mkdir(parents=True, exist_ok=True)
        with open(self._backend_log_path, "a", encoding="utf-8") as fh:
            for line in lines:
                fh.write(line + "\n")

    def _write_record_markdown(self):
        if self._outputs_dir is None:
            return
        self._outputs_dir.mkdir(parents=True, exist_ok=True)
        (self._outputs_dir / RECORD_MARKDOWN_NAME).write_text(
            "# 會議紀錄\n"
            "時間：2026-09-23 10:00\n"
            "## 一、科長轉知\n"
            "1. 宣達事項（00:03:12）。\n"
            "## 二、科長指示及提醒事項\n"
            "1. 提醒事項（00:05:01）。\n",
            encoding="utf-8",
        )

    def post(self, url, files=None, data=None, **kwargs):
        self.upload_calls.append({"url": url, "data": data})
        self._append_log([UPLOAD_LOG_LINE])
        if self._stored_bytes is not None and self._uploads_dir is not None:
            self._uploads_dir.mkdir(parents=True, exist_ok=True)
            (self._uploads_dir / self._stored_filename).write_bytes(self._stored_bytes)
        return FakeResponse(200, self._upload_json)

    def get(self, url, params=None, **kwargs):
        if url.rstrip("/").endswith("/transcript"):
            return self._transcript_response
        if "/result" in url:
            return self._docx_response
        self.task_polls += 1
        self._append_log(self._metrics_lines)
        self._write_record_markdown()
        return FakeResponse(200, self._task_json)


# ---------------------------------------------------------------------------
# DOCX fixtures（python-docx 既有依賴；合成內容）
# ---------------------------------------------------------------------------

def _docx_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _general_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("盤點工具討論 會議紀錄", level=1)
    doc.add_heading("一、報告事項", level=2)
    doc.add_paragraph("資訊室報告系統上線進度，本週已完成驗收測試。")
    doc.add_heading("二、討論事項", level=2)
    doc.add_paragraph("討論年度採購時程、預算配置與分工。")
    doc.add_paragraph("決議：由資訊室於九月底前完成採購文件並送簽。")
    doc.add_heading("三、主席裁示事項", level=2)
    doc.add_paragraph("主席裁示：照案通過，請各單位配合辦理。")
    return _docx_bytes(doc)


def _section_meeting_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("科務會議 會議紀錄", level=1)
    doc.add_heading("一、科長轉知", level=2)
    doc.add_paragraph("轉知公文一件，請同仁依限辦理。")
    doc.add_heading("二、科長指示及提醒事項", level=2)
    doc.add_paragraph("1. 提醒同仁準時出席教育訓練。")
    doc.add_heading("案由及承辦單位", level=2)
    doc.add_paragraph("案由：採購案進度；承辦單位：資訊室。")
    doc.add_heading("散會", level=2)
    doc.add_paragraph("主席宣布散會。")
    return _docx_bytes(doc)


# ---------------------------------------------------------------------------
# 標準 fake flow 組裝
# ---------------------------------------------------------------------------

def _one_model_payload():
    return {
        "models": [
            {
                "key": MODEL_KEY,
                "type": "llm",
                "loaded_instances": [
                    {"id": INSTANCE_ID, "config": {"context_length": 183296}}
                ],
            }
        ]
    }


def _models_unreachable(_call_index):
    return FakeResponse(500, text="LM Studio 不可達（provider 非 lmstudio 的模擬）")


def _default_config_json():
    """預設＝與 Mac LM Studio 現況一致（strict inventory gate）。"""
    return {"local_llm_provider": "auto", "effective_local_llm_provider": "lmstudio"}


def _task_json(*, template_id="general", processing_mode="local", status="completed"):
    return {
        "task_id": TASK_ID,
        "filename": STORED_FILENAME,
        "original_filename": "0903-科務會議.m4a",
        "file_size": len(AUDIO_BYTES),
        "status": status,
        "progress": 100.0,
        "stage": "完成",
        "processing_mode": processing_mode,
        "template_id": template_id,
        "summary_failed": False,
    }


def _upload_json():
    return {
        "task_id": TASK_ID,
        "filename": "0903-科務會議.m4a",
        "file_size": len(AUDIO_BYTES),
        "queue_position": 0,
        "estimated_wait_seconds": 0,
        "message": "檔案上傳成功，已加入處理佇列",
    }


def _transcript_response():
    return FakeResponse(
        200,
        content="會議逐字稿測試內容（00:03:12）宣達事項".encode("utf-8"),
        headers={"content-disposition": 'attachment; filename="20260923120000_逐字稿.txt"'},
    )


def _docx_response(docx_bytes, label=FORMAL_LABEL):
    return FakeResponse(
        200,
        content=docx_bytes,
        headers={"content-disposition": 'attachment; filename="20260923120000_' + label + '.docx"'},
    )


def _record_quality_payload(
    *,
    body=25,
    table=0,
    traceable=0.99,
    on_start=0.98,
    excluding_zero=0.95,
    metric_version=TAG_METRIC_VERSION,
):
    """合成 record 品質量測 payload（schema 對齊 scripts/e2e/measure_record_quality.py）。"""
    return {
        "char_count": 4200,
        "body_source_tag_count": body,
        "table_source_tag_count": table,
        "non_prefixed_tableish_source_tag_count": 0,
        "instruction_item_count": 6,
        "tagged_item_ratio": 0.97,
        "cross_section_duplicate_pairs": 0,
        "known_term_fix_hits": {
            "left_hits": 0,
            "right_hits": 2,
            "left_by_form": {},
            "right_by_form": {},
            "transcript": None,
        },
        "tag_traceability": {
            "tags_total": 30,
            "traceable_tag_ratio": traceable,
            "on_start_tag_ratio": on_start,
            "on_start_tag_ratio_excluding_zero": excluding_zero,
            "metric_version": metric_version,
        },
        "unsupported_entities": ["某某中心"],
        "notes": {},
    }


def _coverage_payload(checklist_sha256):
    return {
        "metric_version": "coverage-1.0.0",
        "label": "runner-coverage-fake",
        "checklist_sha256": checklist_sha256,
        "coverage_all": 0.55,
        "coverage_core": 0.5,
        "fact_total": 67,
        "fact_core_total": 40,
        "covered_total": 37,
        "covered_core": 20,
        "missing_core_ids": ["F001", "F002"],
        "warnings": [],
    }


def _install_fake_instrument(
    monkeypatch,
    *,
    record_payload=None,
    coverage_payload=None,
    record_error=None,
    instrument_exit_code=0,
):
    """取代 runner.run_quality_instrument：只寫 --out／--json-out JSON、不做真 subprocess。"""
    calls = []

    def fake_instrument(script_path, args, *, data_dir, timeout_seconds=None):
        name = Path(str(script_path)).name
        calls.append(
            {
                "script": name,
                "args": [str(arg) for arg in args],
                "data_dir": str(data_dir),
                "timeout_seconds": timeout_seconds,
            }
        )
        if record_error is not None and name == "measure_record_quality.py":
            return {
                "exit_code": None,
                "timed_out": False,
                "stdout": "",
                "stderr": "",
                "error": record_error,
            }
        payload = coverage_payload if name == "measure_coverage.py" else record_payload
        out_flag = "--out" if "--out" in args else "--json-out"
        if out_flag in args:
            out_path = Path(args[args.index(out_flag) + 1])
            out_path.parent.mkdir(parents=True, exist_ok=True)
            out_path.write_text(json.dumps(payload or {}, ensure_ascii=False), encoding="utf-8")
        return {
            "exit_code": instrument_exit_code,
            "timed_out": False,
            "stdout": json.dumps(payload or {}, ensure_ascii=False),
            "stderr": "",
            "error": None,
        }

    monkeypatch.setattr(runner, "run_quality_instrument", fake_instrument)
    return calls


def _run_runner(
    monkeypatch,
    capsys,
    tmp_path,
    *,
    extra_args=(),
    task_json=None,
    config_json=None,
    models_factory=None,
    data_dir=None,
):
    """以 fake stack 執行 runner.main()，回傳可觀察結果。"""
    artifacts_dir = tmp_path / "e2e-artifacts"
    if data_dir is None:
        data_dir = tmp_path / "isolated-data"
    audio_path = tmp_path / "audio.m4a"
    audio_path.write_bytes(AUDIO_BYTES)
    if task_json is None:
        task_json = _task_json()
    docx_bytes = (
        _section_meeting_docx_bytes()
        if "section_meeting" in extra_args
        else _general_docx_bytes()
    )
    client = FakeUploadClient(
        upload_json=_upload_json(),
        task_json=task_json,
        transcript_response=_transcript_response(),
        docx_response=_docx_response(docx_bytes),
        backend_log_path=artifacts_dir / "backend.log",
        uploads_dir=data_dir / "uploads",
        outputs_dir=data_dir / "outputs",
        metrics_lines=[METRICS_OK_LINE],
    )
    fake_proc = FakeChildProc()

    def fake_popen(*args, **kwargs):
        return fake_proc

    shim_subprocess = SimpleNamespace(
        Popen=fake_popen,
        TimeoutExpired=subprocess.TimeoutExpired,
        STDOUT=subprocess.STDOUT,
        PIPE=subprocess.PIPE,
        DEVNULL=subprocess.DEVNULL,
    )
    fake_httpx = FakeHttpxModule(
        health_response=FakeResponse(
            200, {"status": "ok", "version": "test", "build_revision": EXPECTED_REVISION}
        ),
        models_factory=models_factory or (lambda _call: FakeResponse(200, _one_model_payload())),
        config_response=(
            config_json
            if isinstance(config_json, Exception)
            else FakeResponse(200, config_json if config_json is not None else _default_config_json())
        ),
        upload_client=client,
    )
    monkeypatch.setattr(runner, "subprocess", shim_subprocess)
    monkeypatch.setattr(runner, "httpx", fake_httpx)
    monkeypatch.setattr(runner, "git_status_porcelain", lambda: "")
    monkeypatch.setattr(
        sys,
        "argv",
        [
            str(RUNNER_PATH),
            "--artifacts-dir",
            str(artifacts_dir),
            "--expected-revision",
            EXPECTED_REVISION,
            "--audio",
            str(audio_path),
            "--data-dir",
            str(data_dir),
            *extra_args,
        ],
    )

    outcome = {"exit_code": None, "exception": None}
    try:
        outcome["exit_code"] = runner.main()
    except BaseException as exc:  # 含 SystemExit（parser.error）
        outcome["exception"] = exc

    captured = capsys.readouterr()
    summary = None
    summary_path = artifacts_dir / "run_summary.json"
    if summary_path.exists():
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
    return SimpleNamespace(
        exit_code=outcome["exit_code"],
        exception=outcome["exception"],
        summary=summary,
        artifacts_dir=artifacts_dir,
        data_dir=data_dir,
        output=captured.out + captured.err,
        client=client,
        config_calls=fake_httpx.config_calls,
        models_calls=fake_httpx.models_calls,
    )


def _reasons(result) -> list:
    if result.summary is None:
        return []
    return list(result.summary.get("failure_reasons", []))


def _reasons_text(result) -> str:
    return "\n".join(_reasons(result))


def _write_checklist(tmp_path) -> Path:
    checklist = tmp_path / "fact_checklist.json"
    checklist.write_text(
        json.dumps({"checklist_version": "1.0.0", "meeting": "section_meeting", "facts": []}),
        encoding="utf-8",
    )
    return checklist


# ---------------------------------------------------------------------------
# ① off＝函式級 golden：summary 頂層鍵集合／checks 全集／verdict 與現行完全一致
# ---------------------------------------------------------------------------

def test_off_mode_is_function_level_golden(monkeypatch, capsys, tmp_path):
    calls = _install_fake_instrument(monkeypatch, record_payload=_record_quality_payload())
    result = _run_runner(monkeypatch, capsys, tmp_path)

    assert result.exit_code == 0
    assert result.summary is not None
    assert set(result.summary.keys()) == GOLDEN_SUMMARY_KEYS
    assert "quality" not in result.summary and "engine" not in result.summary
    assert result.summary["checks"] == GOLDEN_FULL_CHECKS
    assert result.summary["failure_reasons"] == []
    assert result.summary["verdict"] == "PASS"
    # off 不得觸發品質儀器（函式級 golden；非只跑兩次 live 對照）
    assert calls == []


def test_off_mode_plan_output_has_no_quality_keys(monkeypatch, capsys, tmp_path):
    _install_fake_instrument(monkeypatch, record_payload=_record_quality_payload())
    result = _run_runner(monkeypatch, capsys, tmp_path, extra_args=("--dry-run",))

    assert result.exit_code == 0
    payload = json.loads(result.output[result.output.index("{"):])
    assert "quality_mode" not in payload["plan"]
    assert "coverage_checklist" not in payload["plan"]


# ---------------------------------------------------------------------------
# ② observe：品質值進 run_summary，但 verdict 不因品質值改變
# ---------------------------------------------------------------------------

def test_observe_low_quality_does_not_change_verdict(monkeypatch, capsys, tmp_path):
    low = _record_quality_payload(body=3, table=1, traceable=0.5, on_start=0.4, excluding_zero=0.3)
    _install_fake_instrument(monkeypatch, record_payload=low)
    result = _run_runner(
        monkeypatch,
        capsys,
        tmp_path,
        task_json=_task_json(template_id="section_meeting"),
        extra_args=("--quality-mode", "observe", "--template", "section_meeting"),
    )

    assert result.exit_code == 0
    assert result.summary["verdict"] == "PASS"
    quality = result.summary["quality"]
    assert quality["mode"] == "observe"
    assert quality["captured"] is True
    assert quality["checks"]["quality_body_tag_count_ok"] is False
    assert quality["checks"]["quality_table_tag_count_zero"] is False
    assert quality["metrics"]["body_source_tag_count"] == 3
    assert len(quality["check_failures"]) == 5
    assert "品質" not in _reasons_text(result)
    # engine：mode_used 事後斷言通過、provider 歸屬有值
    assert result.summary["engine"]["mode_used"] == "local"
    assert result.summary["engine"]["mode_used_match"] is True
    assert result.summary["engine"]["effective_local_llm_provider"] == "lmstudio"


# ---------------------------------------------------------------------------
# ③ required：低品質 fixture → FAIL，failure_reasons 逐條指名（既有門檻）
# ---------------------------------------------------------------------------

def test_required_low_quality_fixture_fails_with_named_reasons(monkeypatch, capsys, tmp_path):
    low = _record_quality_payload(body=3, table=1, traceable=0.5, on_start=0.4, excluding_zero=0.3)
    _install_fake_instrument(monkeypatch, record_payload=low)
    result = _run_runner(
        monkeypatch,
        capsys,
        tmp_path,
        task_json=_task_json(template_id="section_meeting"),
        extra_args=("--quality-mode", "required", "--template", "section_meeting"),
    )

    assert result.exit_code == 1
    assert result.summary["verdict"] == "FAIL"
    reasons = _reasons_text(result)
    for fragment in (
        "body_source_tag_count=3",
        "table_source_tag_count=1",
        "traceable_tag_ratio=0.5",
        "on_start_tag_ratio=0.4",
        "on_start_tag_ratio_excluding_zero=0.3",
    ):
        assert fragment in reasons, reasons
    quality = result.summary["quality"]
    assert quality["checks"]["quality_body_tag_count_ok"] is False
    assert result.summary["checks"]["quality_body_tag_count_ok"] is False
    assert result.summary["checks"]["quality_metrics_captured"] is True
    assert result.summary["engine"]["mode_used_match"] is True


def test_required_high_quality_fixture_passes(monkeypatch, capsys, tmp_path):
    _install_fake_instrument(monkeypatch, record_payload=_record_quality_payload())
    result = _run_runner(
        monkeypatch,
        capsys,
        tmp_path,
        task_json=_task_json(template_id="section_meeting"),
        extra_args=("--quality-mode", "required", "--template", "section_meeting"),
    )

    assert result.exit_code == 0
    assert result.summary["verdict"] == "PASS"
    quality = result.summary["quality"]
    assert all(value is True for value in quality["checks"].values())
    assert quality["check_failures"] == []
    assert result.summary["checks"]["quality_table_tag_count_zero"] is True


# ---------------------------------------------------------------------------
# ④ 儀器失敗：observe＝fail-soft（只記錯）；required＝fail-closed（逐條指名）
# ---------------------------------------------------------------------------

def test_observe_instrument_failure_is_fail_soft(monkeypatch, capsys, tmp_path):
    _install_fake_instrument(monkeypatch, record_error="PermissionError: 模擬儀器無法啟動")
    result = _run_runner(
        monkeypatch,
        capsys,
        tmp_path,
        task_json=_task_json(template_id="section_meeting"),
        extra_args=("--quality-mode", "observe", "--template", "section_meeting"),
    )

    assert result.exit_code == 0
    assert result.summary["verdict"] == "PASS"
    quality = result.summary["quality"]
    assert quality["captured"] is False
    assert any("品質儀器失敗" in message for message in quality["errors"])
    assert "品質" not in _reasons_text(result)


def test_required_instrument_failure_fails_closed(monkeypatch, capsys, tmp_path):
    _install_fake_instrument(monkeypatch, record_error="PermissionError: 模擬儀器無法啟動")
    result = _run_runner(
        monkeypatch,
        capsys,
        tmp_path,
        task_json=_task_json(template_id="section_meeting"),
        extra_args=("--quality-mode", "required", "--template", "section_meeting"),
    )

    assert result.exit_code == 1
    assert result.summary["verdict"] == "FAIL"
    reasons = _reasons_text(result)
    assert "品質儀器失敗" in reasons, reasons
    assert "模擬儀器無法啟動" in reasons, reasons
    assert result.summary["checks"]["quality_metrics_captured"] is False


def test_required_instrument_version_mismatch_fails_closed(monkeypatch, capsys, tmp_path):
    _install_fake_instrument(
        monkeypatch,
        record_payload=_record_quality_payload(metric_version="tag_traceability-9.9.9"),
    )
    result = _run_runner(
        monkeypatch,
        capsys,
        tmp_path,
        task_json=_task_json(template_id="section_meeting"),
        extra_args=("--quality-mode", "required", "--template", "section_meeting"),
    )

    assert result.exit_code == 1
    assert "品質儀器版本不符" in _reasons_text(result)


# ---------------------------------------------------------------------------
# ⑤ mode_used 事後斷言（所有 quality-mode 皆生效）
# ---------------------------------------------------------------------------

def test_mode_used_mismatch_fails(monkeypatch, capsys, tmp_path):
    _install_fake_instrument(monkeypatch, record_payload=_record_quality_payload())
    result = _run_runner(
        monkeypatch, capsys, tmp_path, task_json=_task_json(processing_mode="cloud")
    )

    assert result.exit_code == 1
    assert result.summary["verdict"] == "FAIL"
    assert "mode_used 事後斷言失敗" in _reasons_text(result)
    assert "processing_mode='cloud'" in _reasons_text(result)


def test_mode_used_missing_field_fails(monkeypatch, capsys, tmp_path):
    task = _task_json()
    task.pop("processing_mode")
    _install_fake_instrument(monkeypatch, record_payload=_record_quality_payload())
    result = _run_runner(monkeypatch, capsys, tmp_path, task_json=task)

    assert result.exit_code == 1
    assert "mode_used 事後斷言失敗" in _reasons_text(result)


# ---------------------------------------------------------------------------
# ⑥ provider 拆分：非 LM Studio provider 不得因 LM Studio 專屬檢查缺席而 FAIL
# ---------------------------------------------------------------------------

def test_provider_split_exempts_non_lmstudio_provider(monkeypatch, capsys, tmp_path):
    _install_fake_instrument(monkeypatch, record_payload=_record_quality_payload())
    result = _run_runner(
        monkeypatch,
        capsys,
        tmp_path,
        config_json={
            "local_llm_provider": "ollama",
            "effective_local_llm_provider": "ollama",
        },
        models_factory=_models_unreachable,
    )

    assert result.exit_code == 0
    assert result.summary["verdict"] == "PASS"
    checks = result.summary["checks"]
    assert "model_inventory_unique" not in checks
    assert "model_snapshot_consistent" not in checks
    provider_info = json.loads(
        (result.artifacts_dir / "provider_info.json").read_text(encoding="utf-8")
    )
    assert provider_info["lmstudio_inventory_gate"] is False
    assert provider_info["effective_local_llm_provider"] == "ollama"


def test_provider_mismatch_fails(monkeypatch, capsys, tmp_path):
    _install_fake_instrument(monkeypatch, record_payload=_record_quality_payload())
    result = _run_runner(
        monkeypatch,
        capsys,
        tmp_path,
        config_json={
            "local_llm_provider": "lmstudio",
            "effective_local_llm_provider": "ollama",
        },
        models_factory=_models_unreachable,
    )

    assert result.exit_code == 1
    assert result.summary["verdict"] == "FAIL"
    assert "provider 事後斷言失敗" in _reasons_text(result)


def test_provider_probe_unreachable_keeps_strict_gate(monkeypatch, capsys, tmp_path):
    """探測不可達＝provider 未知 → fail-safe 維持現行嚴格 gate（不得誤豁免）。"""
    _install_fake_instrument(monkeypatch, record_payload=_record_quality_payload())
    result = _run_runner(
        monkeypatch,
        capsys,
        tmp_path,
        config_json=RuntimeError("模擬 /api/config 不可達"),
    )

    assert result.exit_code == 0
    assert result.summary["checks"]["model_inventory_unique"] is True
    assert result.summary["checks"]["model_snapshot_consistent"] is True
    provider_info = json.loads(
        (result.artifacts_dir / "provider_info.json").read_text(encoding="utf-8")
    )
    assert provider_info["lmstudio_inventory_gate"] is True
    assert provider_info["error"]


# ---------------------------------------------------------------------------
# ⑦ --coverage-checklist：永遠只做觀測（off 也可給；不影響 verdict）
# ---------------------------------------------------------------------------

def test_coverage_checklist_is_observation_only(monkeypatch, capsys, tmp_path):
    checklist = _write_checklist(tmp_path)
    checklist_sha = hashlib.sha256(checklist.read_bytes()).hexdigest()
    calls = _install_fake_instrument(
        monkeypatch,
        record_payload=_record_quality_payload(),
        coverage_payload=_coverage_payload(checklist_sha),
    )
    result = _run_runner(
        monkeypatch,
        capsys,
        tmp_path,
        extra_args=("--coverage-checklist", str(checklist)),
    )

    assert result.exit_code == 0
    assert result.summary["verdict"] == "PASS"
    coverage = result.summary["quality"]["observations"]["coverage"]
    assert coverage["observation_only"] is True
    assert coverage["gate_effect"] == "none"
    assert coverage["captured"] is True
    assert coverage["metric_version"] == "coverage-1.0.0"
    assert coverage["coverage_core"] == 0.5
    assert coverage["checklist_sha256"] == checklist_sha
    # off 模式仍不寫 engine 鍵；record 品質儀器不執行（僅 coverage 觀測）
    assert "engine" not in result.summary
    assert [call["script"] for call in calls] == ["measure_coverage.py"]
    evidence = json.loads(
        (result.artifacts_dir / "coverage_observation.json").read_text(encoding="utf-8")
    )
    assert evidence["observation_only"] is True and evidence["gate_effect"] == "none"
    assert "quality_gate_passed" not in evidence and "passed" not in evidence


def test_coverage_checklist_failure_does_not_fail_verdict(monkeypatch, capsys, tmp_path):
    _install_fake_instrument(monkeypatch, record_payload=_record_quality_payload())
    missing_checklist = tmp_path / "missing_checklist.json"
    result = _run_runner(
        monkeypatch,
        capsys,
        tmp_path,
        extra_args=("--coverage-checklist", str(missing_checklist)),
    )

    assert result.exit_code == 0
    assert result.summary["verdict"] == "PASS"
    coverage = result.summary["quality"]["observations"]["coverage"]
    assert coverage["captured"] is False
    assert any("checklist 不存在" in message for message in coverage["errors"])


# ---------------------------------------------------------------------------
# ⑧ required 無 --template ＝ parser.error（exit 2）
# ---------------------------------------------------------------------------

def test_required_without_template_is_parser_error(monkeypatch, capsys, tmp_path):
    result = _run_runner(
        monkeypatch, capsys, tmp_path, extra_args=("--quality-mode", "required")
    )

    assert isinstance(result.exception, SystemExit)
    assert result.exception.code == 2
    assert "--quality-mode required 需要 --template" in result.output
    assert result.summary is None


# ---------------------------------------------------------------------------
# ⑨ run_quality_instrument：真 subprocess 的 DATA_DIR／timeout／exit code（fail-soft）
# ---------------------------------------------------------------------------

def test_run_quality_instrument_real_subprocess_semantics(tmp_path):
    data_dir = tmp_path / "probe-data"
    data_dir.mkdir()
    echo_script = tmp_path / "echo_env.py"
    echo_script.write_text(
        "import json, os\n"
        "print(json.dumps({'data_dir': os.environ.get('DATA_DIR'),"
        " 'log_level': os.environ.get('LOG_LEVEL')}))\n",
        encoding="utf-8",
    )
    result = runner.run_quality_instrument(echo_script, [], data_dir=data_dir)
    assert result["exit_code"] == 0 and result["error"] is None
    payload = json.loads(result["stdout"])
    assert payload["data_dir"] == str(data_dir)
    assert payload["log_level"] == "CRITICAL"

    sleep_script = tmp_path / "sleep.py"
    sleep_script.write_text("import time\ntime.sleep(5)\n", encoding="utf-8")
    timed_out = runner.run_quality_instrument(
        sleep_script, [], data_dir=data_dir, timeout_seconds=0.5
    )
    assert timed_out["timed_out"] is True
    assert timed_out["exit_code"] is None
    assert "逾時" in timed_out["error"]

    fail_script = tmp_path / "fail.py"
    fail_script.write_text("import sys\nsys.exit(3)\n", encoding="utf-8")
    failed = runner.run_quality_instrument(fail_script, [], data_dir=data_dir)
    assert failed["exit_code"] == 3
    assert failed["error"] is None


def test_quality_helpers_pure_functions():
    """既有門檻純函式：None 不判定、eq／ge 語意與逐條指名訊息。"""
    metrics = {
        "table_source_tag_count": 0,
        "body_source_tag_count": 20,
        "tag_traceability": {
            "traceable_tag_ratio": 0.99,
            "on_start_tag_ratio": 0.97,
            "on_start_tag_ratio_excluding_zero": None,
        },
    }
    checks, failures, _notes = runner.evaluate_quality_checks(
        metrics, source_tags_applicable=True
    )
    assert failures == []
    assert checks["quality_on_start_excluding_zero_ok"] is None  # None＝不判定
    assert checks["quality_body_tag_count_ok"] is True

    not_applicable, failures, notes = runner.evaluate_quality_checks(
        metrics, source_tags_applicable=False
    )
    assert failures == [] and all(value is None for value in not_applicable.values())
    assert notes

    low, failures, _ = runner.evaluate_quality_checks(
        {"table_source_tag_count": 2, "body_source_tag_count": 1},
        source_tags_applicable=True,
    )
    assert low["quality_table_tag_count_zero"] is False
    assert any("table_source_tag_count=2" in message for message in failures)
    assert any("body_source_tag_count=1" in message for message in failures)


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))
