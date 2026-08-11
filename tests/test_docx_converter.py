"""
DOCX 轉換器測試。

驗證 MarkdownToDocxConverter 的核心功能：
- 標題、段落、粗體、列表、表格、引用、水平線
- CJK 字型設定
- 邊界情況處理
- API 整合（format=docx）
"""

import os
import sys
import tempfile
import pytest

# 確保可匯入專案模組
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.oxml.ns import qn
from backend.services.docx_converter import MarkdownToDocxConverter, FONT_HEADING, FONT_BODY


# ===== 測試用 Markdown 範本 =====

SAMPLE_MEETING_MD = """# 會議記錄摘要

## 1. 會議概況
- **日期**：2026 年 2 月 28 日
- **參與者**：張三、李四、王五
- **會議主題**：AI 自動化專案進度檢討

## 2. 執行摘要
本次會議主要討論 AI 自動化專案導入後的實際執行成效。重點在於準確率提升與使用者接受度。

## 3. 詳細議題與決議
- **議題一：系統準確率評估**
  - 討論重點：目前系統準確率約 85%，仍有提升空間
  - 最終決議：後續需進行更多訓練資料收集

- **議題二：使用者回饋整合**
  - 討論重點：使用者希望增加半自動模式
  - 最終決議：下一版本加入半自動化功能

## 4. 待辦事項
| 待辦事項 | 負責人 | 期限 |
| :--- | :--- | :--- |
| 收集訓練資料 | 張三 | 2026 年 3 月 15 日 |
| 設計半自動介面 | 李四 | 2026 年 3 月 30 日 |
| 撰寫測試報告 | 王五 | 2026 年 4 月 1 日 |

## 5. 其他備註
- 下次會議訂於 3 月 7 日召開
- 需邀請產品經理參與下次討論

---

## 原始逐字稿

<details>
<summary>點擊展開逐字稿</summary>

這是一段測試用的逐字稿內容。
張三說：我們目前的準確率大概在 85% 左右。
李四說：使用者反映希望有半自動模式。

</details>
"""

SAMPLE_SIMPLE_MD = """# 簡單測試

## 標題一
這是一段普通文字。

## 標題二
- 項目一
- 項目二
"""

SAMPLE_TABLE_ONLY_MD = """## 待辦事項
| 事項 | 負責人 | 期限 |
| :--- | :--- | :--- |
| 任務A | 甲 | 明天 |
| 任務B | 乙 | 下週 |
"""

SAMPLE_WITH_BLOCKQUOTE = """# 會議記錄

> 檔案：test_audio.mp3
> 處理時間：2026-02-28 10:00:00
> 處理模式：本地模式 (Ollama)

---

## 1. 會議概況
- **日期**：無
"""


@pytest.fixture
def converter():
    return MarkdownToDocxConverter()


@pytest.fixture
def tmp_docx(tmp_path):
    """回傳暫存 .docx 檔案路徑。"""
    return str(tmp_path / "test_output.docx")


# ===== 基礎轉換測試 =====

class TestBasicConversion:
    """基礎轉換功能測試。"""

    def test_full_meeting_conversion(self, converter, tmp_docx):
        """完整會議記錄 Markdown → DOCX，驗證檔案可正常開啟。"""
        result = converter.convert(SAMPLE_MEETING_MD, tmp_docx)
        assert os.path.exists(result)
        assert result.endswith('.docx')

        # 驗證可用 python-docx 開啟
        doc = Document(tmp_docx)
        assert len(doc.paragraphs) > 0

    def test_simple_conversion(self, converter, tmp_docx):
        """簡單 Markdown 轉換。"""
        converter.convert(SAMPLE_SIMPLE_MD, tmp_docx)
        doc = Document(tmp_docx)
        assert len(doc.paragraphs) > 0

    def test_empty_content_raises(self, converter, tmp_docx):
        """空內容應拋出 ValueError。"""
        with pytest.raises(ValueError, match="不可為空"):
            converter.convert("", tmp_docx)

        with pytest.raises(ValueError, match="不可為空"):
            converter.convert("   ", tmp_docx)

    def test_output_file_created(self, converter, tmp_docx):
        """轉換後檔案應存在。"""
        converter.convert(SAMPLE_SIMPLE_MD, tmp_docx)
        assert os.path.exists(tmp_docx)
        assert os.path.getsize(tmp_docx) > 0


# ===== 標題解析測試 =====

class TestHeadingParsing:
    """標題解析測試。"""

    def test_h1_heading(self, converter, tmp_docx):
        """H1 標題正確解析。"""
        converter.convert("# 測試標題\n\n內文", tmp_docx)
        doc = Document(tmp_docx)

        # 找到 Heading 1 段落
        h1_found = False
        for para in doc.paragraphs:
            if para.style.name.startswith('Heading 1'):
                assert '測試標題' in para.text
                h1_found = True
                break
        assert h1_found, "未找到 H1 標題"

    def test_h2_heading(self, converter, tmp_docx):
        """H2 標題正確解析。"""
        converter.convert("## 二級標題\n\n內文", tmp_docx)
        doc = Document(tmp_docx)

        h2_found = False
        for para in doc.paragraphs:
            if para.style.name.startswith('Heading 2'):
                assert '二級標題' in para.text
                h2_found = True
                break
        assert h2_found, "未找到 H2 標題"

    def test_multiple_headings(self, converter, tmp_docx):
        """多個標題層級同時存在。"""
        md = "# 主標題\n\n## 子標題一\n\n## 子標題二\n"
        converter.convert(md, tmp_docx)
        doc = Document(tmp_docx)

        h1_count = sum(1 for p in doc.paragraphs if p.style.name.startswith('Heading 1'))
        h2_count = sum(1 for p in doc.paragraphs if p.style.name.startswith('Heading 2'))
        assert h1_count == 1
        assert h2_count == 2


# ===== 表格解析測試 =====

class TestTableParsing:
    """表格解析測試 - 重點驗證 Office 2024/M365 相容性。"""

    def test_table_created(self, converter, tmp_docx):
        """Markdown 表格正確轉為 DOCX 表格。"""
        converter.convert(SAMPLE_TABLE_ONLY_MD, tmp_docx)
        doc = Document(tmp_docx)

        assert len(doc.tables) == 1
        table = doc.tables[0]
        # 標題行 + 2 資料行（分隔線行被跳過）
        assert len(table.rows) == 3
        assert len(table.columns) == 3

    def test_table_header_content(self, converter, tmp_docx):
        """表格標題行內容正確。"""
        converter.convert(SAMPLE_TABLE_ONLY_MD, tmp_docx)
        doc = Document(tmp_docx)
        table = doc.tables[0]

        headers = [table.rows[0].cells[i].text.strip() for i in range(3)]
        assert headers == ['事項', '負責人', '期限']

    def test_table_data_content(self, converter, tmp_docx):
        """表格資料行內容正確。"""
        converter.convert(SAMPLE_TABLE_ONLY_MD, tmp_docx)
        doc = Document(tmp_docx)
        table = doc.tables[0]

        # 第一筆資料行
        row1 = [table.rows[1].cells[i].text.strip() for i in range(3)]
        assert row1 == ['任務A', '甲', '明天']

        # 第二筆資料行
        row2 = [table.rows[2].cells[i].text.strip() for i in range(3)]
        assert row2 == ['任務B', '乙', '下週']

    def test_table_has_borders(self, converter, tmp_docx):
        """表格有完整框線（Office 2024/M365 相容）。"""
        converter.convert(SAMPLE_TABLE_ONLY_MD, tmp_docx)
        doc = Document(tmp_docx)
        table = doc.tables[0]

        # 檢查 XML 中有 tblBorders 設定
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.findall(qn('w:tblBorders'))
        assert len(borders) > 0, "表格缺少 tblBorders 設定"

        # 驗證各邊框都有設定
        border_elem = borders[0]
        for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            found = border_elem.findall(qn(f'w:{side}'))
            assert len(found) > 0, f"表格缺少 {side} 框線"

    def test_table_header_has_shading(self, converter, tmp_docx):
        """表格標題行有灰底。"""
        converter.convert(SAMPLE_TABLE_ONLY_MD, tmp_docx)
        doc = Document(tmp_docx)
        table = doc.tables[0]

        # 檢查第一行的儲存格有底色
        header_cell = table.rows[0].cells[0]
        tc_pr = header_cell._tc.tcPr
        shading = tc_pr.findall(qn('w:shd'))
        assert len(shading) > 0, "表格標題行缺少底色"

    def test_full_meeting_table(self, converter, tmp_docx):
        """完整會議記錄中的待辦事項表格。"""
        converter.convert(SAMPLE_MEETING_MD, tmp_docx)
        doc = Document(tmp_docx)

        assert len(doc.tables) >= 1
        table = doc.tables[0]
        # 標題行 + 3 筆資料
        assert len(table.rows) == 4
        assert len(table.columns) == 3

        # 驗證標題
        headers = [table.rows[0].cells[i].text.strip() for i in range(3)]
        assert '待辦事項' in headers[0]
        assert '負責人' in headers[1]
        assert '期限' in headers[2]

    def test_separator_row_skipped(self, converter):
        """表格分隔線行被正確跳過。"""
        assert MarkdownToDocxConverter._is_table_separator("| :--- | :--- | :--- |")
        assert MarkdownToDocxConverter._is_table_separator("| --- | --- | --- |")
        assert MarkdownToDocxConverter._is_table_separator("| :---: | ---: | :--- |")
        assert not MarkdownToDocxConverter._is_table_separator("| 事項 | 負責人 | 期限 |")


# ===== 粗體解析測試 =====

class TestBoldParsing:
    """粗體文字解析測試。"""

    def test_bold_text(self, converter, tmp_docx):
        """**粗體** 正確轉為粗體 Run。"""
        converter.convert("## 測試\n\n- **日期**：2026 年\n", tmp_docx)
        doc = Document(tmp_docx)

        bold_found = False
        for para in doc.paragraphs:
            for run in para.runs:
                if run.bold and '日期' in run.text:
                    bold_found = True
                    break
        assert bold_found, "未找到粗體文字"

    def test_mixed_bold_normal(self, converter, tmp_docx):
        """粗體與一般文字混合。"""
        converter.convert("## 測試\n\n這是 **重要** 的文字\n", tmp_docx)
        doc = Document(tmp_docx)

        has_bold = False
        has_normal = False
        for para in doc.paragraphs:
            for run in para.runs:
                if run.bold:
                    has_bold = True
                else:
                    has_normal = True
        assert has_bold, "缺少粗體 Run"
        assert has_normal, "缺少一般 Run"

    def test_same_text_bold_and_plain_not_confused(self, converter, tmp_docx):
        """同一行有相同文字的粗體與非粗體片段時，位置不得前後誤置（v4.6.1）。"""
        converter.convert("## 測試\n\n重要**重要**\n", tmp_docx)
        doc = Document(tmp_docx)

        target = next(p for p in doc.paragraphs if p.text == "重要重要")
        assert target.runs[0].text == "重要"
        assert target.runs[0].bold is not True  # 前段為一般文字
        assert target.runs[1].text == "重要"
        assert target.runs[1].bold is True  # 後段才是 **粗體**


# ===== 列表項目測試 =====

class TestListItems:
    """列表項目解析測試。"""

    def test_bullet_list(self, converter, tmp_docx):
        """一級項目符號。"""
        md = "## 測試\n\n- 項目一\n- 項目二\n- 項目三\n"
        converter.convert(md, tmp_docx)
        doc = Document(tmp_docx)

        bullet_items = [p for p in doc.paragraphs if p.style.name == 'List Bullet']
        assert len(bullet_items) == 3

    def test_nested_bullet_list(self, converter, tmp_docx):
        """二級項目符號。"""
        md = "## 測試\n\n- 一級項目\n  - 二級項目 A\n  - 二級項目 B\n"
        converter.convert(md, tmp_docx)
        doc = Document(tmp_docx)

        l1_items = [p for p in doc.paragraphs if p.style.name == 'List Bullet']
        l2_items = [p for p in doc.paragraphs if p.style.name == 'List Bullet 2']
        assert len(l1_items) == 1
        assert len(l2_items) == 2


# ===== 引用區塊測試 =====

class TestBlockquote:
    """引用區塊測試。"""

    def test_blockquote_italic(self, converter, tmp_docx):
        """引用區塊應為斜體。"""
        converter.convert(SAMPLE_WITH_BLOCKQUOTE, tmp_docx)
        doc = Document(tmp_docx)

        italic_found = False
        for para in doc.paragraphs:
            for run in para.runs:
                if run.italic and '檔案' in run.text:
                    italic_found = True
                    break
        assert italic_found, "引用區塊未設定為斜體"


# ===== CJK 字型測試 =====

class TestCJKFonts:
    """CJK 字型設定測試。"""

    def test_default_font_set(self, converter, tmp_docx):
        """文件預設字型應為新細明體。"""
        converter.convert(SAMPLE_SIMPLE_MD, tmp_docx)
        doc = Document(tmp_docx)

        normal_style = doc.styles['Normal']
        assert normal_style.font.name == FONT_BODY

    def test_east_asian_font_set(self, converter, tmp_docx):
        """East Asian 字型屬性應設定正確。"""
        converter.convert(SAMPLE_SIMPLE_MD, tmp_docx)
        doc = Document(tmp_docx)

        normal_style = doc.styles['Normal']
        east_asian = normal_style.element.rPr.rFonts.get(qn('w:eastAsia'))
        assert east_asian == FONT_BODY

    def test_heading_font(self, converter, tmp_docx):
        """標題字型應為微軟正黑體。"""
        converter.convert("# 測試標題\n\n內文", tmp_docx)
        doc = Document(tmp_docx)

        for para in doc.paragraphs:
            if para.style.name.startswith('Heading'):
                for run in para.runs:
                    assert run.font.name == FONT_HEADING
                    break


# ===== 頁面設定測試 =====

class TestPageSetup:
    """頁面設定測試。"""

    def test_a4_page_size(self, converter, tmp_docx):
        """頁面應為 A4 大小。"""
        converter.convert(SAMPLE_SIMPLE_MD, tmp_docx)
        doc = Document(tmp_docx)

        section = doc.sections[0]
        # A4: 21cm x 29.7cm (允許微小誤差)
        assert abs(section.page_width.cm - 21.0) < 0.1
        assert abs(section.page_height.cm - 29.7) < 0.1


# ===== details/summary HTML 標籤測試 =====

class TestHtmlTags:
    """HTML 標籤處理測試。"""

    def test_details_tag_stripped(self, converter, tmp_docx):
        """<details> 標籤應被移除，內容保留。"""
        converter.convert(SAMPLE_MEETING_MD, tmp_docx)
        doc = Document(tmp_docx)

        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert '<details>' not in full_text
        assert '<summary>' not in full_text
        assert '</details>' not in full_text
        # 但逐字稿內容應保留
        assert '測試用的逐字稿' in full_text


# ===== 水平線測試 =====

class TestHorizontalRule:
    """水平分隔線測試。"""

    def test_horizontal_rule_present(self, converter, tmp_docx):
        """--- 應轉為段落分隔線。"""
        converter.convert(SAMPLE_WITH_BLOCKQUOTE, tmp_docx)
        doc = Document(tmp_docx)

        # 驗證文件中有段落底部框線（水平線實作方式）
        hr_found = False
        for para in doc.paragraphs:
            pPr = para._p.find(qn('w:pPr'))
            if pPr is not None:
                pBdr = pPr.find(qn('w:pBdr'))
                if pBdr is not None:
                    hr_found = True
                    break
        assert hr_found, "未找到水平分隔線"


# ===== 端到端整合測試 =====

class TestEndToEnd:
    """端到端整合測試：模擬完整會議記錄轉換流程。"""

    def test_full_pipeline_content_integrity(self, converter, tmp_docx):
        """完整轉換後，所有關鍵內容都應保留。"""
        converter.convert(SAMPLE_MEETING_MD, tmp_docx)
        doc = Document(tmp_docx)

        full_text = "\n".join(p.text for p in doc.paragraphs)

        # 標題
        assert '會議記錄摘要' in full_text
        # 會議概況
        assert '2026 年 2 月 28 日' in full_text
        assert '張三' in full_text
        # 執行摘要
        assert 'AI 自動化專案' in full_text
        # 議題
        assert '系統準確率評估' in full_text
        # 備註
        assert '3 月 7 日' in full_text

    def test_full_pipeline_table_integrity(self, converter, tmp_docx):
        """完整轉換後，待辦事項表格資料完整。"""
        converter.convert(SAMPLE_MEETING_MD, tmp_docx)
        doc = Document(tmp_docx)

        table = doc.tables[0]
        all_text = ""
        for row in table.rows:
            for cell in row.cells:
                all_text += cell.text + " "

        assert '收集訓練資料' in all_text
        assert '張三' in all_text
        assert '設計半自動介面' in all_text
        assert '李四' in all_text
        assert '撰寫測試報告' in all_text
        assert '王五' in all_text

    def test_file_size_reasonable(self, converter, tmp_docx):
        """輸出檔案大小合理（不應過大或過小）。"""
        converter.convert(SAMPLE_MEETING_MD, tmp_docx)
        size = os.path.getsize(tmp_docx)
        # DOCX 至少有基本結構 (> 5KB)，不應超過 1MB（純文字內容）
        assert size > 5000, f"檔案過小: {size} bytes"
        assert size < 1_000_000, f"檔案過大: {size} bytes"


# =============================================================================
# v4.3.2：公文紀錄結構樣式（純文字紀錄不再是同大小字牆）
# =============================================================================

class TestRecordStructureStyling:
    def test_section_and_label_lines_are_styled(self, tmp_path):
        """「一、報告事項」章節應加粗放大；「案由：/決議：」標籤應加粗。"""
        md = (
            "# 會議紀錄\n\n"
            "會議名稱：測試會議\n"
            "主  席：局長\n\n"
            "一、 報告事項：\n"
            "無。\n\n"
            "二、 討論事項：\n"
            "案由：測試案由內容。\n"
            "決議：照案通過。\n"
        )
        out = str(tmp_path / "record.docx")
        MarkdownToDocxConverter().convert(md, out)

        doc = Document(out)
        by_text = {p.text: p for p in doc.paragraphs if p.text}

        section = by_text["一、 報告事項："]
        assert section.runs[0].bold is True
        assert section.runs[0].font.size.pt == 15

        label = by_text["案由：測試案由內容。"]
        assert label.runs[0].text == "案由："
        assert label.runs[0].bold is True
        assert label.runs[1].bold is not True  # 內容不加粗

        meta = by_text["會議名稱：測試會議"]
        assert meta.runs[0].text == "會議名稱："
        assert meta.runs[0].bold is True

        # 全形空白欄位（主  席）也要能辨識
        chair = by_text["主  席：局長"]
        assert chair.runs[0].bold is True


# =============================================================================
# v4.6.1：章節白名單——雲端紀錄的編號條列項目不得整段粗體放大
# （2026-08-11 使用者回報：雲端模式 DOCX 標題與內文全成粗體）
# =============================================================================

class TestNumberedItemsRenderAsBody:
    """模擬雲端模型輸出：意見／決議以「一、二、三、」條列。"""

    CLOUD_STYLE_MD = (
        "# 會議紀錄\n\n"
        "會議名稱：測試會議\n\n"
        "一、 報告事項：\n"
        "無。\n\n"
        "二、 討論事項：\n"
        "案由一：AI 工具功能調整，提請 審議。\n"
        "說明：背景說明。\n"
        "各單位意見（多方立場）：\n"
        "一、 科長表示，本案應由人工檢視後判定。\n"
        "二、 經討論後，決議捨棄複雜分類。\n"
        "決議：\n"
        "一、 取消自動判定功能，交由人工確認。\n"
        "二、 檔案版本不強制使用者手動填寫，欄位保留空白。\n\n"
        "三、 主席裁示事項（後續管考與追蹤）：\n"
        "一、 請重新提供業務分類碼清單。\n"
    )

    def _convert(self, tmp_path):
        out = str(tmp_path / "cloud_record.docx")
        MarkdownToDocxConverter().convert(self.CLOUD_STYLE_MD, out)
        doc = Document(out)
        return {p.text: p for p in doc.paragraphs if p.text}

    def test_real_sections_keep_heading_style(self, tmp_path):
        by_text = self._convert(tmp_path)
        for section_text in (
            "一、 報告事項：",
            "二、 討論事項：",
            "三、 主席裁示事項（後續管考與追蹤）：",
        ):
            run = by_text[section_text].runs[0]
            assert run.bold is True, f"章節應加粗: {section_text}"
            assert run.font.size.pt == 15

    def test_numbered_items_render_as_plain_body(self, tmp_path):
        by_text = self._convert(tmp_path)
        for item_text in (
            "一、 科長表示，本案應由人工檢視後判定。",
            "二、 經討論後，決議捨棄複雜分類。",
            "一、 取消自動判定功能，交由人工確認。",
            "二、 檔案版本不強制使用者手動填寫，欄位保留空白。",
            "一、 請重新提供業務分類碼清單。",
        ):
            for run in by_text[item_text].runs:
                assert run.bold is not True, f"條列項目不得粗體: {item_text}"
                assert run.font.size is None, f"條列項目不得放大: {item_text}"
                assert run.font.name == FONT_BODY

    def test_numbered_proposal_label_bold(self, tmp_path):
        """「案由一：」標籤加粗、內容照常。"""
        by_text = self._convert(tmp_path)
        proposal = by_text["案由一：AI 工具功能調整，提請 審議。"]
        assert proposal.runs[0].text == "案由一："
        assert proposal.runs[0].bold is True
        assert proposal.runs[1].bold is not True


# =============================================================================
# v4.6.0：ISMS 月工作會議表單式版面（FormLayoutSpec 忠實還原官方會議記錄表）
# =============================================================================

ISMS_FORM_MD = """機關名稱：嘉義縣財政稅務局
專案名稱：嘉義縣財政稅務局115年度資通安全管理與個人資料保護制度整合委外維護案
會議議題：嘉義縣財政稅務局115年度資通安全管理與個人資料保護制度整合委外維護案
7月工作執行報告
地點：電子作業科會議室
主席：何科長
日期：115.07.28（二）
記錄：龔君
機關單位：財政稅務局電子作業科
廠商單位：天雷科技有限公司
參加人員：
嘉義縣財政稅務局：陳股長、王管理師及李助理稅務員
駐點工程師：張工程師
天雷科技顧問：龔君
內容：
一、天雷科技有限公司針對7月份工作進行說明。
二、天雷科技有限公司針對8月份工作安排說明。
追蹤事項：
無
決議事項：
一、與AI智慧分文之承辦人員討論「人工智慧應用風險處理計畫表」規劃改善事項及預計完成時間。
二、房屋稅清查所識別之人工智慧應用風險，決議接受該風險。
三、內部稽核時間為9月7日；外部驗證時間為10月14日。
四、下次工作會議為8月20日上午10:00。
臨時動議：
無。"""


class TestIsmsFormDocx:
    """ISMS 表單式 DOCX 渲染測試（Document() 開回逐項驗證版面結構）。"""

    def _convert(self, md, tmp_docx):
        MarkdownToDocxConverter().convert(md, tmp_docx, template_id="isms_monthly")
        return Document(tmp_docx)

    def test_two_tables_with_expected_rows(self, tmp_docx):
        doc = self._convert(ISMS_FORM_MD, tmp_docx)
        assert len(doc.tables) == 2
        assert len(doc.tables[0].rows) == 8   # 主表格
        assert len(doc.tables[1].rows) == 3   # 簽到表

    def test_intro_paragraphs(self, tmp_docx):
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

        doc = self._convert(ISMS_FORM_MD, tmp_docx)
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        assert texts[0] == "嘉義縣財政稅務局 會議記錄表"
        assert texts[1] == "□專案啟動會議；■月工作會議"
        title = next(p for p in doc.paragraphs if "會議記錄表" in p.text)
        assert title.alignment == WD_PARAGRAPH_ALIGNMENT.CENTER
        assert title.runs[0].bold is True

    def test_main_table_field_placement(self, tmp_docx):
        doc = self._convert(ISMS_FORM_MD, tmp_docx)
        main = doc.tables[0]
        assert main.rows[0].cells[0].text.strip() == "專案名稱"
        assert "整合委外維護案" in main.rows[0].cells[-1].text
        topic_cell = main.rows[1].cells[0]
        assert "會議議題：" in topic_cell.text
        assert "7月工作執行報告" in topic_cell.text
        assert main.rows[2].cells[0].text.strip() == "地    點"
        assert "電子作業科會議室" in main.rows[2].cells[1].text
        assert "何科長" in main.rows[2].cells[-1].text
        assert "115.07.28（二）" in main.rows[3].cells[1].text
        assert "龔君" in main.rows[3].cells[-1].text

    def test_decision_and_motion_share_cell(self, tmp_docx):
        """決議事項與臨時動議照範本同一儲存格，決議逐條各自成段。"""
        doc = self._convert(ISMS_FORM_MD, tmp_docx)
        cell = doc.tables[0].rows[7].cells[0]
        assert "決議事項：" in cell.text
        assert "臨時動議：" in cell.text
        assert "決議接受該風險" in cell.text
        # 標籤2段＋決議4條＋臨時動議1行＋區塊間空段 ≥ 7 段
        assert len(cell.paragraphs) >= 7

    def test_sign_in_table(self, tmp_docx):
        doc = self._convert(ISMS_FORM_MD, tmp_docx)
        sign = doc.tables[1]
        assert "簽" in sign.rows[0].cells[0].text and "到" in sign.rows[0].cells[0].text
        assert "財政稅務局電子作業科" in sign.rows[1].cells[0].text
        assert "天雷科技有限公司" in sign.rows[1].cells[1].text
        # 簽名列為空格且有最小列高
        assert sign.rows[2].cells[0].text.strip() == ""
        tr_height = sign.rows[2]._tr.trPr.find(qn('w:trHeight'))
        assert tr_height is not None
        assert tr_height.get(qn('w:hRule')) == "atLeast"

    def test_table_xml_office_compat(self, tmp_docx):
        """框線與固定版面須以明確 XML 設定（Office 2024/M365 相容）。"""
        doc = self._convert(ISMS_FORM_MD, tmp_docx)
        for table in doc.tables:
            tbl_pr = table._tbl.tblPr
            assert tbl_pr.find(qn('w:tblBorders')) is not None
            layout = tbl_pr.find(qn('w:tblLayout'))
            assert layout is not None and layout.get(qn('w:type')) == "fixed"

    def test_missing_fields_render_confirm_placeholder(self, tmp_docx):
        """欄位缺漏時表格仍完整產出，缺格填（待確認）。"""
        doc = self._convert("專案名稱：某案\n決議事項：\n一、某決議。", tmp_docx)
        main = doc.tables[0]
        assert "（待確認）" in main.rows[2].cells[1].text  # 地點
        assert "（待確認）" in main.rows[6].cells[0].text  # 追蹤事項
        assert "一、某決議。" in main.rows[7].cells[0].text

    def test_transcript_appendix_appended_after_tables(self, tmp_docx):
        """水平線後的原始逐字稿附錄須以一般渲染接在表格之後，不遺失內容。"""
        md = f"{ISMS_FORM_MD}\n\n---\n\n## 原始逐字稿\n\n這是逐字稿內容。"
        doc = self._convert(md, tmp_docx)
        assert len(doc.tables) == 2
        texts = [p.text for p in doc.paragraphs]
        assert any("這是逐字稿內容" in t for t in texts)
        assert "逐字稿" not in doc.tables[0].rows[7].cells[0].text

    def test_fallback_to_general_rendering_without_form_labels(self, tmp_docx):
        """本文未命中適用樣式（如摘要失敗為逐字稿）→ 退回一般渲染不拋例外。"""
        doc = self._convert("# 標題\n\n這是一段沒有表單欄位標籤的內容。", tmp_docx)
        assert len(doc.tables) == 0
        assert any("沒有表單欄位標籤" in p.text for p in doc.paragraphs)

    def test_other_templates_unaffected(self, tmp_docx):
        """一般模板照舊走段落式渲染（回歸保證）。"""
        MarkdownToDocxConverter().convert(SAMPLE_MEETING_MD, tmp_docx, template_id="general")
        doc = Document(tmp_docx)
        assert len(doc.paragraphs) > 0
